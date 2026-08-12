import asyncio
from pathlib import Path

from hashmm.agent.office_artifacts import SCHEMA, inspect_office, report_markdown
from hashmm.api.design_quality import audit_html
from hashmm.evaluation.rag_diagnostics import diagnose_run
import pytest


def test_docx_inspection_is_deterministic_and_structural(tmp_path: Path):
    from docx import Document

    target = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("真实标题", level=1)
    doc.add_paragraph("这是程序实际读取的正文。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(1, 0).text = "值"
    doc.save(target)

    report = inspect_office(target)
    assert report["schema"] == SCHEMA
    assert report["verified"] is True
    assert report["structure"]["heading_count"] == 1
    assert report["structure"]["table_count"] == 1
    assert len(report["file"]["sha256"]) == 64
    assert "真实标题" in report_markdown(report)


def test_pptx_and_xlsx_inspection_reads_real_structure(tmp_path: Path):
    from openpyxl import Workbook
    from pptx import Presentation

    deck = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "第一页"
    slide.placeholders[1].text = "要点"
    prs.save(deck)
    assert inspect_office(deck)["structure"]["slides"][0]["title"] == "第一页"

    book = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "收入"
    ws.append(["项目", "金额"])
    ws.append(["A", 10])
    ws["C2"] = "=B2*2"
    wb.save(book)
    report = inspect_office(book)
    assert report["structure"]["sheets"][0]["name"] == "收入"
    assert report["structure"]["sheets"][0]["formulas_scanned"] == 1


def test_office_tool_is_chat_wired_and_rejects_traversal(tmp_path: Path, monkeypatch):
    from hashmm.agent import modules, permissions
    from hashmm.api import tool_registry as registry

    monkeypatch.setattr(registry, "get_files_dir", lambda _conv=None: tmp_path)
    denied = registry._exec_inspect_office({"filename": "../foreign.docx"}, {"conv_id": "mine"})
    assert denied["status"] == "error"
    assert "不能包含路径" in denied["message"]
    assert registry.get_executor("inspect_office") is not None
    assert "inspect_office" in {t["function"]["name"] for t in registry.TOOL_DEFS}
    assert permissions.TOOL_PERMISSIONS["inspect_office"] == permissions.PermissionLevel.READ
    computer = next(item for item in modules.all_modules() if item.key == "computer")
    assert "inspect_office" in computer.tools


def test_design_gate_blocks_remote_dependencies_and_records_accessibility_warnings():
    unsafe = audit_html("<html><head><script src='https://example.test/a.js'></script></head><body><h1>x</h1></body></html>")
    assert unsafe["passed"] is False
    assert unsafe["checks"]["offline_self_contained"] is False

    local = audit_html("<html lang='zh'><head><meta name='viewport' content='width=device-width'><title>x</title></head><body><img src='data:image/png;base64,AA'></body></html>")
    assert local["passed"] is True
    assert "1 张图片缺少 alt" in local["warnings"]


def test_rag_diagnostics_only_claims_observable_patterns():
    result = diagnose_run(
        sources=[{"text": "短片段"}, {"text": "又一个"}, {"text": "第三个"}],
        groundings={"total_factual_claims": 4, "supported_claims": 1,
                    "coverage_ratio": 0.25, "review_required": True,
                    "semantic_entailment_verified": False},
        corpus={"status": "evidence_only"},
        iterations=12,
        tool_steps=[{"tool": "kb_search", "status": "failed"}],
    )
    ids = {item["id"] for item in result["observed"]}
    assert {"P01_grounding_drift", "P02_chunk_boundary", "P06_long_chain_drift",
            "P07_tool_reliability", "P09_eval_blind_spot",
            "P11_config_reproducibility"} <= ids
    unevaluated = {item["id"] for item in result["not_evaluable"]}
    assert "P03_embedding_mismatch" in unevaluated
    assert result["integrity"]["causes_inferred_without_signal"] is False


def test_office_upload_is_validated_before_replacing_existing_file(tmp_path: Path):
    from fastapi import HTTPException
    from hashmm.api.routes.conversations import _write_conversation_upload

    target = tmp_path / "report.docx"
    original = b"known-good-placeholder"
    target.write_bytes(original)
    with pytest.raises(HTTPException) as raised:
        asyncio.run(_write_conversation_upload(target, b"not-an-office-archive"))
    assert raised.value.status_code == 422
    assert target.read_bytes() == original
    assert not list(tmp_path.glob(".*.upload-*.docx"))
