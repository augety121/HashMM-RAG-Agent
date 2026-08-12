"""Conversation routes — CRUD, files, export, execute, fork, tags.

v17 Phase 63 (security): every /conversations/{conv_id} endpoint now goes
through require_conv_access(...) before touching data. Previously these
endpoints operated purely by id with no ownership check, so anyone who knew a
conv_id could read/modify/delete/export another user's conversation or run code
in their workspace (OWASP API #1 — Broken Object Level Authorization). The
guard returns the conversation row (admins bypass ownership) or raises 404.
"""
from __future__ import annotations
from hashmm.utils import get_logger as _get_logger, log_suppressed
_obs_logger = _get_logger(__name__)
import base64, hashlib, json, math, os, re, time, uuid
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from pathlib import Path
from fastapi import APIRouter, Request, HTTPException, UploadFile, File
from fastapi.responses import FileResponse, PlainTextResponse, Response, HTMLResponse, JSONResponse
from pydantic import BaseModel, Field
from starlette.concurrency import run_in_threadpool

from hashmm.api import database as db
from hashmm.api import app_state
from hashmm.api.auth import (
    get_current_user, require_auth, require_conv_access, require_conv_access_or_create, resolve_user_id,
)
# Phase 63 fix: execute_tool was used below (/execute, /files/create) but never
# imported — those endpoints raised NameError at call time. Import from
# tool_registry (its source module) NOT server: server imports the routes
# package, so importing from server here would create a circular import.
from hashmm.api.tool_registry import execute_tool
from hashmm.pipeline.resource_pipeline import parse_resource, resource_summary

router = APIRouter(prefix="/api", tags=["conversations"])


_SECRET_QUERY_KEY = re.compile(r"(?:token|key|auth|session|password|passwd|secret|credential|code)", re.I)


def _decode_conversation_cursor(
    value: str,
    *,
    scope: str,
    project_id: str | None,
) -> tuple[int, float, str] | None:
    if not value:
        return None
    try:
        padded = value + "=" * (-len(value) % 4)
        payload = json.loads(base64.urlsafe_b64decode(padded.encode("ascii")))
        if payload.get("v") != 1:
            raise ValueError("unsupported cursor")
        if str(payload.get("s") or "all") != scope:
            raise ValueError("cursor scope mismatch")
        if str(payload.get("project") or "") != str(project_id or ""):
            raise ValueError("cursor project mismatch")
        pinned = int(payload["p"])
        activity = float(payload["u"])
        conv_id = str(payload["i"])
        if pinned not in (0, 1) or not conv_id or not math.isfinite(activity):
            raise ValueError("invalid cursor values")
        return pinned, activity, conv_id
    except (KeyError, TypeError, ValueError, json.JSONDecodeError, UnicodeError):
        raise HTTPException(status_code=422, detail="invalid_conversation_cursor")


def _encode_conversation_cursor(
    row: dict,
    *,
    scope: str,
    project_id: str | None,
) -> str:
    payload = {
        "v": 1,
        "s": scope,
        "project": str(project_id or ""),
        "p": 1 if row.get("pinned") else 0,
        "u": float(row.get("last_activity_at") or row.get("updated_at") or 0),
        "i": str(row.get("id") or ""),
    }
    raw = json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8")
    return base64.urlsafe_b64encode(raw).decode("ascii").rstrip("=")


def _bounded_text(value, limit: int, *, one_line: bool = False) -> str:
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]+", "", str(value or ""))
    if one_line:
        text = re.sub(r"\s+", " ", text).strip()
    return text[: max(0, limit)]


def _safe_external_url(value) -> str:
    """Remove credentials, fragments and secret-like query parameters."""
    try:
        parsed = urlsplit(str(value or ""))
        if parsed.scheme not in ("http", "https") or not parsed.hostname:
            return ""
        host = parsed.hostname
        if ":" in host and not host.startswith("["):
            host = f"[{host}]"
        if parsed.port:
            host = f"{host}:{parsed.port}"
        query = urlencode([(key, val) for key, val in parse_qsl(parsed.query, keep_blank_values=True)
                           if not _SECRET_QUERY_KEY.search(key)])
        return urlunsplit((parsed.scheme, host, parsed.path, query, ""))[:500]
    except (TypeError, ValueError):
        return ""


def _external_tool_trace(value) -> list[dict]:
    """Whitelist the bounded trajectory shape accepted from desktop clients."""
    out: list[dict] = []
    for raw in (value[:40] if isinstance(value, list) else []):
        if not isinstance(raw, dict):
            continue
        status = str(raw.get("status") or raw.get("state") or "done")
        if status == "failed":
            status = "error"
        if status not in ("running", "done", "error"):
            status = "done"
        item = {
            "id": _bounded_text(raw.get("id"), 96, one_line=True) or f"external-{len(out) + 1}",
            "node": _bounded_text(raw.get("node") or raw.get("tool"), 96, one_line=True) or "desktop:step",
            "tool": _bounded_text(raw.get("tool"), 48, one_line=True),
            "detail": _bounded_text(raw.get("detail"), 420, one_line=True),
            "status": status,
        }
        try:
            elapsed = int(raw.get("elapsed_ms"))
            if 0 <= elapsed <= 86_400_000:
                item["elapsed_ms"] = elapsed
        except (TypeError, ValueError):
            pass
        out.append(item)
    return out


def _external_evidence_contract(content: str, sources) -> tuple[list, dict]:
    """Normalize desktop/browser evidence and build the same Chat ledger.

    External desktop pages are untrusted and may send oversized records.  Only
    the bounded public source contract is persisted; the deterministic ledger
    never upgrades an uncited claim merely because a browser source exists.
    """
    if not isinstance(sources, list) or not sources:
        return [], {}
    from hashmm.evaluation.grounding_ledger import build_grounding_ledger, public_sources
    safe: list[dict] = []
    seen_browser_urls: set[str] = set()
    for raw in sources[:10]:
        if not isinstance(raw, dict):
            continue
        evidence_text = _bounded_text(raw.get("text") or raw.get("snippet") or raw.get("content"), 1200)
        if not evidence_text.strip():
            continue
        method = _bounded_text(raw.get("method") or raw.get("via"), 32, one_line=True)
        section = _bounded_text(raw.get("section") or raw.get("loc"), 500, one_line=True)
        if method == "browser_use":
            section = _safe_external_url(section)
            if not section or section in seen_browser_urls:
                continue
            seen_browser_urls.add(section)
        try:
            score = float(raw.get("score", 0.0))
            if not math.isfinite(score):
                score = 0.0
        except (TypeError, ValueError):
            score = 0.0
        safe.append({
            "citation_id": len(safe) + 1,
            "source_id": _bounded_text(raw.get("source_id") or raw.get("id"), 128, one_line=True)
                         or f"external-source-{len(safe) + 1}",
            "chunk_id": _bounded_text(raw.get("chunk_id"), 128, one_line=True),
            "doc_id": _bounded_text(raw.get("doc_id"), 128, one_line=True),
            "filename": _bounded_text(raw.get("filename") or raw.get("file") or raw.get("source"), 160, one_line=True)
                        or "外部来源",
            "page": raw.get("page", -1),
            "section": section,
            "score": max(-1.0, min(1.0, score)),
            "method": method,
            "modality": _bounded_text(raw.get("modality") or "text", 32, one_line=True),
            "text": evidence_text,
        })
    if not safe:
        return [], {}
    return public_sources(safe), build_grounding_ledger(content, safe)


class TitleRequest(BaseModel):
    query: str
    answer: str = ""


class TurnSteerRequest(BaseModel):
    content: str
    client_message_id: str = ""
    attachments: list[dict] = Field(default_factory=list)


def _verified_steer_attachments(conv_id: str, values: list[dict]) -> list[dict]:
    """Accept only files already committed to this owned conversation.

    Client-provided paths and URLs are ignored.  The server recomputes the
    receipt from bytes inside the owner-checked conversation workspace.
    """
    if len(values or []) > 8:
        raise HTTPException(422, "一次最多追加 8 个附件")
    root = db.conv_files_dir(conv_id).resolve()
    out: list[dict] = []
    for raw in values or []:
        if not isinstance(raw, dict):
            raise HTTPException(422, "附件信息格式错误")
        name = str(raw.get("filename") or "").strip()
        safe = Path(name).name
        if not safe or safe != name or safe in (".", ".."):
            raise HTTPException(422, "附件名称不合法")
        path = (root / safe).resolve()
        if path.parent != root or not path.is_file():
            raise HTTPException(422, "附件不在当前会话工作区")
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        claimed = str(raw.get("sha256") or "").lower()
        if claimed and claimed != digest:
            raise HTTPException(409, "附件在上传后发生了变化，请重新添加")
        out.append({
            "filename": safe,
            "size": len(data),
            "sha256": digest,
            "download_url": f"/api/conversations/{conv_id}/download/{safe}",
        })
    return out


# V50: 标题生成的云端硬超时（秒）。超过即回退本地截取。
_TITLE_TIMEOUT_S = 8.0


@router.post("/title")
async def generate_title(req: TitleRequest):
    """Generate a short conversation title from query+answer.

    The LLM call is synchronous and can take many seconds (slow reasoning
    models). Running it directly in this async handler would block the whole
    event loop — starving every other concurrent request (observed: a 29s title
    call stalling unrelated /corpus/stats requests). So we offload it to a
    thread, keeping the loop responsive.
    """
    llm = app_state.llm_fn
    if not llm:
        return {"title": req.query[:30]}
    import asyncio
    # Cloud/local routing: title generation is a cheap, high-volume task → prefer
    # the local model when routing is enabled (cost + latency). Falls back to the
    # cloud fn transparently.
    try:
        from hashmm.llm_router import route_llm, record_routing
        routed, backend = route_llm("title", llm, user_id=getattr(req, "user_id", None))
        llm = routed or llm
        record_routing("title", backend)
    except Exception:
        pass  # nosem: observability-fallback
    try:
        prompt = (f"为以下对话生成一个简短标题（10字以内，不要引号）：\n"
                  f"问：{req.query[:100]}\n答：{req.answer[:100]}\n标题：")
        # Titles are tiny; cap local generation hard so a 7B model doesn't spend
        # seconds emitting tokens we'll throw away (observed 15s otherwise).
        def _call(p):
            try:
                return llm(p, max_new_tokens=24)  # local fn supports the kwarg
            except TypeError:
                return llm(p)                      # cloud fn: no kwarg
        # V50: 云端推理模型在这种 5 字小活上可能跑 19s+（真机日志 SLOW 19441ms）。
        # 超时直接回退本地截取——标题不值得让用户等；后台线程跑完即作废，无副作用。
        t = await asyncio.wait_for(asyncio.to_thread(_call, prompt), timeout=_TITLE_TIMEOUT_S)
        title = t.strip().strip('"\'「」').split('\n')[0][:30]
        return {"title": title if title else req.query[:30]}
    except asyncio.TimeoutError:
        return {"title": req.query[:30]}
    except Exception:
        return {"title": req.query[:30]}



@router.get("/conversations/{conv_id}/transaction")
async def get_transaction(conv_id: str, request: Request):
    """V300 第二期：某会话/任务的事务日志——Agent 做过哪些有副作用的写操作（可回放/审计）。
    返回 [{step, op, target, ok, detail, created}]。"""
    require_conv_access(request, conv_id)
    try:
        from hashmm.agent import idempotency as _idem
        return {"conv_id": conv_id, "steps": _idem.get_transaction(conv_id)}
    except Exception:
        return {"conv_id": conv_id, "steps": []}


@router.get("/conversations")
async def list_convs(
    request: Request,
    since: str = "",
    cloud_sync: int = 0,
    archived: int = 0,
    limit: int = 200,
    before: float | None = None,
    offset: int = 0,
    include_empty: int = 0,
    scope: str = "all",
    project_id: str | None = None,
    cursor: str = "",
):
    """Return the owner-scoped server-local conversation snapshot.

    Normal reads never contact or write Supabase.  ``cloud_sync=1`` is an
    explicit recovery operation for an old cloud-only account; the blocking
    compatibility pull is isolated from the asyncio event loop.  Mutations are
    mirrored through the write path, so a list request can never create the
    read -> backfill -> Realtime -> read feedback loop.
    """
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    clean_scope = str(scope or "all").strip().lower()
    if clean_scope not in {"all", "unassigned", "project"}:
        raise HTTPException(status_code=422, detail="invalid_conversation_scope")
    clean_project_id = str(project_id or "").strip() or None
    if clean_scope == "project":
        if not clean_project_id or not db.get_project_for_user(clean_project_id, user_id):
            # Missing and unauthorized projects deliberately share one result.
            raise HTTPException(status_code=404, detail="project_not_found")
    elif clean_project_id:
        raise HTTPException(status_code=422, detail="project_id_requires_project_scope")
    cursor_key = _decode_conversation_cursor(
        cursor, scope=clean_scope, project_id=clean_project_id,
    )
    if cursor_key is not None and (offset or before is not None):
        raise HTTPException(status_code=422, detail="cursor_cannot_be_combined_with_offset")
    from datetime import datetime, timezone
    # Capture before the remote query. A racing cloud write is either included
    # now or remains newer than this cursor for the next incremental pull.
    sync_cursor = datetime.now(timezone.utc).isoformat()
    page_limit = max(1, min(int(limit or 200), 500))
    page_offset = max(0, int(offset or 0))
    # Read one extra row so clients can page without guessing whether an exact
    # page was the end of the owner's history.
    convs_page = db.list_conversations(
        user_id, limit=page_limit + 1, before=before, archived=archived,
        offset=page_offset, meaningful_only=not bool(include_empty),
        scope=clean_scope, project_id=clean_project_id, cursor=cursor_key,
    )
    has_more = len(convs_page) > page_limit
    convs = convs_page[:page_limit]
    if cloud_sync == 1:
      try:
        from hashmm.api import supabase_sync
        # V243: 归档区查询(archived != 0)纯本地过滤——跳过云端补建。
        # 此前补建逻辑会把云端会话(select 不含 archived)当"缺失"重新 create_conversation
        # (archived 默认 0)，污染归档列表并使"还原后再打开仍是全部"。归档是本地视图，不掺云同步。
        if archived == 0 and page_offset == 0 and cursor_key is None and supabase_sync.enabled():
            remote = await run_in_threadpool(
                supabase_sync.pull_conversations, user_id, since or None,
            )
            changed = False
            for sc in remote:
                changed = db.merge_cloud_conversation(sc, user_id) or changed
            if changed:
                convs_page = db.list_conversations(
                    user_id, limit=page_limit + 1, before=before,
                    archived=archived, offset=page_offset,
                    meaningful_only=not bool(include_empty),
                    scope=clean_scope, project_id=clean_project_id, cursor=cursor_key,
                )
                has_more = len(convs_page) > page_limit
                convs = convs_page[:page_limit]
      except Exception as exc:
        log_suppressed(_obs_logger, exc)
    return {
        "conversations": convs,
        "sync_cursor": sync_cursor,
        "page": {
            "limit": page_limit,
            "has_more": has_more,
            "next_offset": page_offset + len(convs) if has_more else None,
            "next_cursor": _encode_conversation_cursor(
                convs[-1], scope=clean_scope, project_id=clean_project_id,
            ) if has_more and convs else None,
        },
    }


@router.get("/conversations/search")
async def search_convs(
    request: Request,
    q: str,
    scope: str = "all",
    project_id: str | None = None,
    limit: int = 50,
    cursor: str = "",
):
    """Search every owned Chat while preserving its one navigation home."""
    query = str(q or "").strip()[:240]
    if not query:
        raise HTTPException(status_code=422, detail="search_query_required")
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    clean_scope = str(scope or "all").strip().lower()
    if clean_scope not in {"all", "unassigned", "project"}:
        raise HTTPException(status_code=422, detail="invalid_conversation_scope")
    clean_project_id = str(project_id or "").strip() or None
    if clean_scope == "project":
        if not clean_project_id or not db.get_project_for_user(clean_project_id, user_id):
            raise HTTPException(status_code=404, detail="project_not_found")
    elif clean_project_id:
        raise HTTPException(status_code=422, detail="project_id_requires_project_scope")
    cursor_scope = "search:" + hashlib.sha256(query.encode("utf-8")).hexdigest()[:16] + ":" + clean_scope
    cursor_key = _decode_conversation_cursor(
        cursor, scope=cursor_scope, project_id=clean_project_id,
    )
    page_limit = max(1, min(int(limit or 50), 100))
    rows = db.list_conversations(
        user_id,
        limit=page_limit + 1,
        archived=0,
        meaningful_only=True,
        scope=clean_scope,
        project_id=clean_project_id,
        cursor=cursor_key,
        search_query=query,
    )
    has_more = len(rows) > page_limit
    items = rows[:page_limit]
    project_names: dict[str, str] = {}
    for item in items:
        pid = str(item.get("project_id") or "")
        if pid and pid not in project_names:
            project = db.get_project_for_user(pid, user_id)
            project_names[pid] = str((project or {}).get("name") or "")
        item["location"] = {
            "kind": "project" if pid else "unassigned",
            "project_id": pid or None,
            "project_name": project_names.get(pid, "") if pid else "",
        }
    return {
        "query": query,
        "conversations": items,
        "page": {
            "limit": page_limit,
            "has_more": has_more,
            "next_cursor": _encode_conversation_cursor(
                items[-1], scope=cursor_scope, project_id=clean_project_id,
            ) if has_more and items else None,
        },
    }


@router.get("/conversations/tombstones")
async def get_conversation_tombstones(
    request: Request,
    since: float = 0,
    after_id: str = "",
    limit: int = 500,
):
    """Return owner-scoped deletion markers for offline reconciliation."""
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    page_size = max(1, min(2000, int(limit or 500)))
    rows = db.list_conversation_tombstones(
        user_id, since=since, after_id=after_id, limit=page_size + 1,
    )
    has_more = len(rows) > page_size
    items = rows[:page_size]
    last = items[-1] if items else None
    return {
        "schema": "hashmm.conversation-tombstones.v1",
        "tombstones": items,
        "next_since": float(last["deleted_at"]) if last else float(since or 0),
        "next_after_id": str(last["conversation_id"]) if last else str(after_id or ""),
        "has_more": has_more,
    }


@router.get("/activity")
async def get_activity(request: Request):
    """当前用户客户端正在进行的任务：流式生成中的对话 + 运行中的后台任务(RAG解析/索引)。

    供 App 显示「客户端任务进度」并一键接管(打开对应会话——内容随客户端流式实时同步)。
    流式对话按用户过滤；后台任务为系统级，仅管理员可见，避免多用户泄露。
    """
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    active_chats = db.get_active_chats(user_id)
    jobs_running = []
    if user and user.get("role") == "admin":
        try:
            from hashmm.api import jobs as _jobs
            jobs_running = [
                {"id": j.get("id"), "kind": j.get("kind"), "status": j.get("status"),
                 "done": j.get("done", 0), "total": j.get("total", 0), "message": j.get("message", "")}
                for j in _jobs.list_jobs(limit=50)
                if j.get("status") in ("pending", "running")
            ]
        except Exception:
            jobs_running = []
    return {"active_chats": active_chats, "jobs": jobs_running}


@router.post("/conversations")
async def create_conv(request: Request):
    """Create a new conversation."""
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    body = await request.json()
    conv_id = body.get("id") or str(uuid.uuid4())
    title = body.get("title", "新对话")
    project_id = str(body.get("project_id") or "").strip() or None
    if project_id and not db.get_project_for_user(project_id, user_id):
        raise HTTPException(status_code=404, detail="项目不存在")
    conv = db.create_conversation(
        conv_id, user_id, title, project_id=project_id,
    )
    return {
        "id": conv_id,
        "title": title,
        "project_id": conv.get("project_id"),
    }


@router.post("/conversations/{source_conv_id}/handoffs")
async def create_chat_handoff(source_conv_id: str, request: Request):
    """Seal public task state into a new, owner-scoped continuation Chat."""
    source = require_conv_access(request, source_conv_id)
    owner_id = str(source.get("user_id") or "")
    idempotency_key = str(request.headers.get("idempotency-key") or "").strip()
    if len(idempotency_key) < 8 or len(idempotency_key) > 200:
        raise HTTPException(status_code=422, detail="idempotency_key_required")
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        from hashmm.api import chat_handoffs
        handoff, created = chat_handoffs.create(
            owner_id,
            source_conv_id,
            idempotency_key=idempotency_key,
            target_title=str(body.get("target_title") or ""),
            expires_in_seconds=int(body.get("expires_in_seconds") or 7 * 86400),
        )
    except LookupError:
        raise HTTPException(status_code=404, detail="conversation_not_found")
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:160])
    return JSONResponse(
        status_code=201 if created else 200,
        content={"handoff": handoff, "created": created},
    )


@router.get("/conversations/{target_conv_id}/handoffs/incoming")
async def get_incoming_chat_handoff(target_conv_id: str, request: Request):
    target = require_conv_access(request, target_conv_id)
    from hashmm.api import chat_handoffs
    handoff = chat_handoffs.incoming(str(target.get("user_id") or ""), target_conv_id)
    return {"handoff": handoff}


@router.post("/conversations/{target_conv_id}/handoffs/{handoff_id}/{action}")
async def transition_chat_handoff(
    target_conv_id: str,
    handoff_id: str,
    action: str,
    request: Request,
):
    target = require_conv_access(request, target_conv_id)
    owner_id = str(target.get("user_id") or "")
    current = db.get_conversation_handoff(handoff_id, owner_id)
    if not current or str(current.get("target_conversation_id") or "") != target_conv_id:
        raise HTTPException(status_code=404, detail="handoff_not_found")
    if action not in {"claim", "acknowledge", "reject", "supersede"}:
        raise HTTPException(status_code=422, detail="invalid_handoff_action")
    try:
        updated = db.transition_conversation_handoff(handoff_id, owner_id, action)
    except RuntimeError as exc:
        raise HTTPException(status_code=409, detail=str(exc)[:160])
    if not updated:
        raise HTTPException(status_code=404, detail="handoff_not_found")
    from hashmm.api import chat_handoffs
    updated["state_verification"] = chat_handoffs.verify_current_state(updated)
    return {"handoff": updated}


@router.post("/conversations/{source_conv_id}/mailbox")
async def create_chat_mailbox_message(source_conv_id: str, request: Request):
    """Send a bounded notice to another Chat; notices never carry authority."""
    source = require_conv_access(request, source_conv_id)
    owner_id = str(source.get("user_id") or "")
    idempotency_key = str(request.headers.get("idempotency-key") or "").strip()
    if len(idempotency_key) < 8 or len(idempotency_key) > 200:
        raise HTTPException(status_code=422, detail="idempotency_key_required")
    try:
        body = await request.json()
    except Exception:
        body = {}
    target_conv_id = str(body.get("target_conversation_id") or "").strip()
    if not target_conv_id:
        raise HTTPException(status_code=422, detail="target_conversation_id_required")
    try:
        item, created = db.create_conversation_mailbox_message(
            owner_id=owner_id,
            source_conversation_id=source_conv_id,
            target_conversation_id=target_conv_id,
            payload=body.get("payload") if isinstance(body.get("payload"), dict) else {},
            idempotency_key=idempotency_key,
            expires_at=time.time() + max(60, min(int(body.get("expires_in_seconds") or 86400), 30 * 86400)),
        )
    except LookupError:
        raise HTTPException(status_code=404, detail="conversation_not_found")
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:160])
    return JSONResponse(status_code=201 if created else 200, content={"message": item, "created": created})


@router.get("/conversations/{target_conv_id}/mailbox")
async def get_chat_mailbox(target_conv_id: str, request: Request, limit: int = 50):
    target = require_conv_access(request, target_conv_id)
    owner_id = str(target.get("user_id") or "")
    return {"messages": db.list_conversation_mailbox(owner_id, target_conv_id, limit=limit)}


@router.post("/conversations/{target_conv_id}/mailbox/{message_id}/{action}")
async def transition_chat_mailbox(
    target_conv_id: str,
    message_id: str,
    action: str,
    request: Request,
):
    target = require_conv_access(request, target_conv_id)
    owner_id = str(target.get("user_id") or "")
    current = db.get_conversation_mailbox(message_id, owner_id)
    if not current or str(current.get("target_conversation_id") or "") != target_conv_id:
        raise HTTPException(status_code=404, detail="mailbox_message_not_found")
    if action not in {"read", "acknowledge", "reject"}:
        raise HTTPException(status_code=422, detail="invalid_mailbox_action")
    try:
        updated = db.transition_conversation_mailbox(message_id, owner_id, action)
    except RuntimeError as exc:
        raise HTTPException(status_code=409, detail=str(exc)[:160])
    return {"message": updated}


@router.get("/conversations/{conv_id}")
async def get_conv(conv_id: str, request: Request, cloud_sync: int = 0):
    """Get a local snapshot; cloud recovery is explicit and off-loop."""
    conv = require_conv_access(request, conv_id)
    page_limit = 500
    msgs = db.get_latest_messages(conv_id, limit=page_limit)
    if cloud_sync == 1:
      try:
        from hashmm.api import supabase_sync
        if supabase_sync.enabled():
            uid = conv.get("user_id")
            remote = await run_in_threadpool(supabase_sync.pull_messages, conv_id, uid)
            if remote and len(remote) > len(msgs):
                db.import_messages_local(conv_id, remote)
                msgs = db.get_latest_messages(conv_id, limit=page_limit)
      except Exception as exc:
        log_suppressed(_obs_logger, exc)
    total = db.count_messages(conv_id)
    return {"conversation": conv, "messages": msgs,
            "page": {"total": total, "has_more": total > len(msgs),
                     "oldest_created_at": msgs[0].get("created_at") if msgs else None}}


@router.delete("/conversations/{conv_id}")
async def delete_conv(conv_id: str, request: Request):
    conv = require_conv_access(request, conv_id)
    if not db.delete_conversation(conv_id, owner_id=str(conv.get("user_id") or "")):
        raise HTTPException(status_code=404, detail="conversation_not_found")
    return {"ok": True}


@router.patch("/conversations/{conv_id}")
async def update_conv(conv_id: str, request: Request):
    # V246 归档还原根治：归档列表里很多会话是 App 建的、只在云端、本地 SQLite 没有。
    # require_conv_access 对本地缺失会话直接 404 → 归档/还原 PATCH 静默失败(前端 catch 吞掉)
    # → 状态没落库 → 还原后再打开仍是全部。改用"缺失就补建"，让归档操作对云端会话也生效。
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    body = await request.json()
    existing = db.get_conversation(conv_id)
    if existing is None:
        # A missing cache row is not ownership evidence. Recover only through
        # an owner-scoped cloud query; never let a caller claim an arbitrary id.
        try:
            from hashmm.api import supabase_sync
            remote = (
                await run_in_threadpool(supabase_sync.pull_conversation, conv_id, user_id)
                if supabase_sync.enabled() else None
            )
            if remote:
                db.merge_cloud_conversation(remote, user_id)
                existing = db.get_conversation(conv_id)
        except Exception:
            existing = None
        if existing is None:
            raise HTTPException(status_code=404, detail="对话不存在")
        require_conv_access(request, conv_id)
    else:
        # 本地有：正常鉴权（非属主/管理员 → 404）
        require_conv_access(request, conv_id)
    db.update_conversation(conv_id, **body)
    # 同步到云端（含 archived），保证多端一致
    try:
        from hashmm.api import supabase_sync
        if supabase_sync.enabled():
            row = db.get_conversation(conv_id)
            if row:
                supabase_sync.push_conversation(row)
    except Exception:
        pass
    return {"ok": True}


@router.get("/conversations/{conv_id}/messages")
async def get_conv_messages(
    conv_id: str,
    request: Request,
    limit: int = 100,
    cloud_sync: int = 0,
    before_ts: float | None = None,
):
    """Return the fast server-local message snapshot.

    This endpoint is polled while a remote task runs, so its normal read path
    must not wait on Supabase. Server message writes are already pushed by the
    database layer. ``?cloud_sync=1`` remains available for an explicit,
    one-shot recovery of cloud-only rows.
    """
    conv = require_conv_access(request, conv_id)
    msgs = db.get_latest_messages(conv_id, limit=limit, before_ts=before_ts)
    if cloud_sync == 1:
        try:
            from hashmm.api import supabase_sync
            if supabase_sync.enabled():
                uid = conv.get("user_id") if conv else None
                remote = await run_in_threadpool(supabase_sync.pull_messages, conv_id, uid)
                if remote and len(remote) > len(msgs):
                    db.import_messages_local(conv_id, remote)
                    msgs = db.get_latest_messages(conv_id, limit=limit, before_ts=before_ts)
        except Exception:
            # Cloud recovery is best-effort; the local snapshot remains valid.
            pass
    total = db.count_messages(conv_id)
    remaining_scope = db.count_messages_before(conv_id, before_ts) if before_ts is not None else total
    return {"messages": msgs,
            "page": {"total": total,
                     "has_more": remaining_scope > len(msgs),
                     "oldest_created_at": msgs[0].get("created_at") if msgs else None}}


@router.post("/conversations/{conv_id}/compact")
async def compact_conversation_context(conv_id: str, request: Request):
    """Create a durable manual checkpoint without deleting full messages."""
    require_conv_access(request, conv_id)
    from hashmm.agent.conv_compact import prepare_persistent_history

    prepared = prepare_persistent_history(
        db, conv_id, force=True, trigger="manual",
    )
    state = db.get_context_compaction(conv_id) or {}
    return {
        "ok": True,
        "compacted": bool(prepared.compacted_now),
        "degraded": bool(getattr(prepared, "degraded", False)),
        "degraded_reason": str(getattr(prepared, "degraded_reason", "") or "")[:240],
        "reason": "" if prepared.compacted_now else "新增消息不足，当前检查点已是最新",
        "state": {
            "source_messages": int(state.get("source_messages") or 0),
            "estimated_tokens": int(state.get("estimated_tokens") or prepared.working_tokens),
            "compaction_count": int(state.get("compaction_count") or 0),
            "through_rowid": int(
                state.get("through_rowid") or getattr(prepared, "through_rowid", 0) or 0
            ),
            "checkpoint_schema": str(
                getattr(prepared, "checkpoint_schema", "hashmm.context-checkpoint.v3")
            ),
            "compaction": dict(getattr(prepared, "compaction", {}) or {}),
            "last_trigger": str(state.get("last_trigger") or ""),
            "updated_at": float(state.get("updated_at") or 0),
        },
    }


@router.get("/conversations/{conv_id}/active-turn")
async def get_active_turn(conv_id: str, request: Request):
    """Return only the caller's currently executing turn for this conversation."""
    require_conv_access(request, conv_id)
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    from hashmm.api.active_runs import registry

    run = registry.maybe_get(conv_id, user_id)
    return {"active": bool(run), "turn": run.public() if run else None}


@router.post("/conversations/{conv_id}/turns/{turn_id}/steer")
async def steer_active_turn(
    conv_id: str, turn_id: str, body: TurnSteerRequest, request: Request,
):
    """Append a real user instruction to one exact running Agent turn.

    The expected turn id makes delayed desktop/App requests safe: a steer for a
    previous turn cannot accidentally change a newer task.
    """
    require_conv_access(request, conv_id)
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    username = user["sub"] if user else "anonymous"
    from hashmm.api.active_runs import (
        ActiveRunInvalidSteer, ActiveRunNotFound, ActiveRunNotSteerable,
        ActiveRunQueueFull, registry,
    )

    try:
        run = registry.get(conv_id, user_id, turn_id)
        attachments = _verified_steer_attachments(conv_id, body.attachments)
        entry, duplicate = run.enqueue_steer(
            body.content, body.client_message_id, attachments
        )
    except ActiveRunNotFound as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ActiveRunInvalidSteer as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except (ActiveRunNotSteerable, ActiveRunQueueFull) as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc

    if duplicate:
        return {
            "accepted": True,
            "duplicate": True,
            "turn_id": run.turn_id,
            "message_id": entry.message_id,
        }

    try:
        if entry.attachments:
            message_id = db.create_message(
                conv_id, "user", entry.content,
                files=[dict(item) for item in entry.attachments],
            )
        else:
            message_id = db.create_message(conv_id, "user", entry.content)
        run.bind_message(entry.entry_id, message_id)
        db.audit(user_id, username, "turn_steer",
                 f"turn={turn_id[:24]} chars={len(entry.content)} "
                 f"attachments={len(entry.attachments)}")
    except Exception:
        # A steer is acknowledged only when its durable user message exists.
        run.rollback_steer(entry.entry_id)
        raise
    return {
        "accepted": True,
        "duplicate": False,
        "turn_id": run.turn_id,
        "message_id": message_id,
        "queued_at": entry.queued_at,
    }


@router.post("/conversations/{conv_id}/turns/{turn_id}/interrupt")
async def interrupt_active_turn(conv_id: str, turn_id: str, request: Request):
    """Request cooperative cancellation of one exact running turn."""
    require_conv_access(request, conv_id)
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    username = user["sub"] if user else "anonymous"
    from hashmm.api.active_runs import ActiveRunNotFound, registry

    try:
        run = registry.get(conv_id, user_id, turn_id)
    except ActiveRunNotFound:
        return {
            "accepted": False,
            "turn_id": turn_id,
            "status": "already_finished",
        }
    accepted = run.request_interrupt()
    if not accepted:
        return {
            "accepted": False,
            "turn_id": run.turn_id,
            "status": "already_finished",
        }
    db.audit(user_id, username, "turn_interrupt", f"turn={turn_id[:24]}")
    return {"accepted": True, "turn_id": run.turn_id, "status": "interrupting"}


@router.patch("/conversations/{conv_id}/messages/{msg_id}")
async def update_msg(conv_id: str, msg_id: str, request: Request):
    """Update a streaming message (intermediate save)."""
    require_conv_access(request, conv_id)
    body = await request.json()
    if isinstance(body.get("sources"), list):
        body["sources"], body["groundings"] = _external_evidence_contract(
            str(body.get("content") or ""), body.get("sources"))
    if isinstance(body.get("tool_calls"), list):
        body["tool_calls"] = _external_tool_trace(body.get("tool_calls"))
    db.update_message(msg_id, **body)
    return {"ok": True}


@router.post("/conversations/{conv_id}/tool-approvals/{request_id}")
async def decide_tool_approval(conv_id: str, request_id: str, request: Request):
    """Approve or decline one exact, durable Agent tool invocation.

    This endpoint only records the human decision. It never executes the tool;
    execution remains inside the normal guarded Agent loop, where the stored
    fingerprint is atomically consumed once.
    """
    require_conv_access(request, conv_id)
    user = get_current_user(request)
    actor_id = user["uid"] if user else "anonymous"
    body = await request.json()
    decision = str(body.get("decision") or "").strip().lower()
    if decision not in ("approve", "decline"):
        raise HTTPException(status_code=400, detail="decision 必须是 approve 或 decline")
    row = db.decide_tool_approval(
        request_id=request_id,
        conv_id=conv_id,
        actor_user_id=actor_id,
        approve=decision == "approve",
        actor_is_admin=bool(user and user.get("role") == "admin"),
    )
    if not row:
        # Missing and unauthorized requests deliberately share one response.
        raise HTTPException(status_code=404, detail="审批请求不存在")
    public = db.public_tool_approval(row)
    # The decision is durable run state, not merely a message decoration.
    # Replays from another device converge through the idempotency key and can
    # never approve a different run/call pair.
    work_run_id = str(row.get("work_run_id") or "")
    if work_run_id:
        try:
            from hashmm.agent import work_runtime
            work_runtime.append_event_once(
                work_run_id,
                user_id=str(row.get("user_id") or actor_id),
                event_type="approval_decision",
                status="waiting_approval" if decision == "approve" else "blocked",
                summary=(
                    "用户已批准原工具调用，等待执行器精确续接"
                    if decision == "approve"
                    else "用户已拒绝工具调用，任务保持阻塞"
                ),
                payload={
                    "approval_id": request_id,
                    "decision": decision,
                    "step_id": str(row.get("step_id") or ""),
                    "call_id": str(row.get("call_id") or ""),
                    "tool_name": str(row.get("tool_name") or "")[:96],
                },
                idempotency_key=f"approval-decision:{request_id}:{decision}",
            )
        except Exception as exc:
            log_suppressed(_obs_logger, exc)
    db.audit(actor_id, str((user or {}).get("username") or actor_id),
             "tool_approval_decision",
             f"request={request_id[:24]} tool={row.get('tool_name','')[:80]} decision={decision}")
    return {"ok": True, "approval_request": public}


@router.get("/conversations/{conv_id}/files")
async def list_conv_files(conv_id: str, request: Request):
    require_conv_access(request, conv_id)
    files = db.list_conversation_files(conv_id)
    # Also check filesystem for files not yet in DB
    fdir = db.conv_files_dir(conv_id)
    if fdir.exists():
        for f in sorted(fdir.iterdir()):
            if f.is_file() and not any(cf["filename"] == f.name for cf in files):
                st = f.stat()
                files.append({"filename": f.name, "path": str(f), "size": st.st_size,
                              "mtime": st.st_mtime,
                              "download_url": f"/api/conversations/{conv_id}/download/{f.name}"})
    for f in files:
        f["download_url"] = f"/api/conversations/{conv_id}/download/{f['filename']}"
        f["size_str"] = f"{f.get('size',0)/1024:.1f}K" if f.get("size",0) > 1024 else f"{f.get('size',0)}B"
        # 补充生成时间（文件实际 mtime），供前端显示正确日期，而非 DB 里可能过期的时间
        if not f.get("mtime"):
            try:
                _fp = fdir / f["filename"]
                if _fp.exists():
                    f["mtime"] = _fp.stat().st_mtime
            except Exception:
                pass
    public_files = []
    for item in files:
        # Database paths are server implementation details and may contain a
        # username or workspace root. Never expose them to desktop/App clients.
        filename = str(item.get("filename") or "")
        if not filename or Path(filename).name != filename or filename in (".", ".."):
            continue
        size = int(item.get("size") or 0)
        mtime = item.get("mtime") or item.get("created_at") or 0
        try:
            file_path = fdir / filename
            if file_path.exists() and file_path.is_file():
                file_stat = file_path.stat()
                size = int(file_stat.st_size)
                mtime = float(file_stat.st_mtime)
        except Exception:
            pass
        public_files.append({
            "id": item.get("id"),
            "filename": filename,
            "size": size,
            "size_str": f"{size/1024:.1f}K" if size > 1024 else f"{size}B",
            "mime_type": item.get("mime_type") or "",
            "created_at": item.get("created_at"),
            "mtime": mtime,
            "download_url": f"/api/conversations/{conv_id}/download/{filename}",
        })
    public_files.sort(key=lambda entry: float(entry.get("mtime") or 0), reverse=True)
    revision = hashlib.sha256(json.dumps([
        (entry["filename"], entry["size"], int(float(entry.get("mtime") or 0) * 1000))
        for entry in public_files
    ], ensure_ascii=False, separators=(",", ":")).encode("utf-8")).hexdigest()
    etag = f'"files-{revision}"'
    response_headers = {"ETag": etag, "Cache-Control": "private, no-cache"}
    if request.headers.get("if-none-match") == etag:
        return Response(status_code=304, headers=response_headers)
    return JSONResponse({"files": public_files, "revision": revision}, headers=response_headers)


@router.get("/conversations/{conv_id}/files/{filename}")
async def get_conv_file(conv_id: str, filename: str, request: Request):
    return await _serve_conv_file(conv_id, filename, request)


@router.get("/conversations/{conv_id}/download/{filename}")
async def download_conv_file(conv_id: str, filename: str, request: Request):
    """专用下载路由（路径段 /download/ 与 /files/ 区分，避免任何路由匹配歧义）。"""
    return await _serve_conv_file(conv_id, filename, request)


async def _serve_conv_file(conv_id: str, filename: str, request: Request):
    import urllib.parse, os.path as _osp, mimetypes
    decoded = _osp.basename(urllib.parse.unquote(filename).split("?")[0].split("#")[0])
    # 必须先做属主校验，再解析工作区或枚举文件；否则未授权请求会通过日志和
    # 时序差异探测目录内容。文件名只使用 basename 解码结果，禁止 ../ 穿越。
    try:
        require_conv_access(request, conv_id)
    except HTTPException as _e:
        _obs_logger.warning(f"[下载拒绝] conv={conv_id} status={_e.status_code}")
        raise
    fdir = db.conv_files_dir(conv_id)
    fp = fdir / decoded
    if fp.exists() and fp.is_file():
        inline = request.query_params.get("inline") == "1"
        response_headers = {}
        try:
            if fp.suffix.lower() in _WRITE_EXT_ALLOW and fp.stat().st_size <= _WRITE_SIZE_CAP:
                response_headers["X-HashMM-Content-Sha256"] = hashlib.sha256(fp.read_bytes()).hexdigest()
        except OSError as exc:
            log_suppressed(_obs_logger, exc)
        return FileResponse(
            fp,
            filename=decoded,
            media_type=mimetypes.guess_type(decoded)[0] or "application/octet-stream",
            content_disposition_type="inline" if inline else "attachment",
            headers=response_headers,
        )
    try:
        if fdir.exists():
            targets = {decoded.strip()}
            for p in fdir.iterdir():
                if p.is_file() and p.name.strip() in targets:
                    inline = request.query_params.get("inline") == "1"
                    response_headers = {}
                    try:
                        if p.suffix.lower() in _WRITE_EXT_ALLOW and p.stat().st_size <= _WRITE_SIZE_CAP:
                            response_headers["X-HashMM-Content-Sha256"] = hashlib.sha256(p.read_bytes()).hexdigest()
                    except OSError as exc:
                        log_suppressed(_obs_logger, exc)
                    return FileResponse(
                        p,
                        filename=p.name,
                        media_type=mimetypes.guess_type(p.name)[0] or "application/octet-stream",
                        content_disposition_type="inline" if inline else "attachment",
                        headers=response_headers,
                    )
    except Exception:
        pass
    _obs_logger.error(f"[下载诊断] 文件查找失败：'{filename}' 不在 {fdir}")
    raise HTTPException(404, "文件不存在")



@router.get("/conversations/{conv_id}/export")
async def export_conversation(conv_id: str, request: Request, format: str = "markdown"):
    """Export conversation as Markdown or JSON."""
    conv = require_conv_access(request, conv_id)
    msgs = db.get_messages(conv_id, limit=500)
    files = db.list_conversation_files(conv_id)

    if format == "json":
        return {"conversation": conv, "messages": msgs, "files": files}

    # Markdown export
    lines = [f"# {conv.get('title', '对话')}\n"]
    lines.append(f"*导出时间: {time.strftime('%Y-%m-%d %H:%M')}*\n\n---\n")
    for m in msgs:
        role = "用户" if m["role"] == "user" else "助手"
        lines.append(f"\n### {role}\n\n{m.get('content', '')}\n")
        if m.get("files"):
            for f in (m["files"] if isinstance(m["files"], list) else []):
                fname = f.get("filename", "") if isinstance(f, dict) else str(f)
                if fname:
                    lines.append(f"\n[附件] {fname}\n")
        lines.append("\n---\n")

    md_content = "\n".join(lines)
    return PlainTextResponse(md_content, media_type="text/markdown",
                             headers={"Content-Disposition": f"attachment; filename={conv_id}.md"})



@router.post("/conversations/{conv_id}/upload-zip")
async def upload_zip_endpoint(conv_id: str, request: Request, file: UploadFile = File(...)):
    """Upload and extract ZIP to conversation workspace."""
    require_conv_access(request, conv_id)
    from hashmm.api.workspace import upload_zip
    import tempfile
    # Save upload to temp file
    with tempfile.NamedTemporaryFile(delete=False, suffix=".zip") as tmp:
        content = await file.read()
        tmp.write(content)
        tmp_path = tmp.name
    result = upload_zip(tmp_path, conv_id)
    os.unlink(tmp_path)
    return result


@router.get("/conversations/{conv_id}/tree")
async def get_file_tree(conv_id: str, request: Request):
    """Get file tree for conversation workspace."""
    require_conv_access(request, conv_id)
    from hashmm.api.workspace import file_tree, scan_project
    return {
        "tree_text": file_tree(conv_id),
        "project": scan_project(conv_id),
    }


@router.get("/conversations/{conv_id}/search")
async def search_workspace(conv_id: str, request: Request, q: str = ""):
    """Search text in workspace files."""
    require_conv_access(request, conv_id)
    from hashmm.api.workspace import search_in_workspace
    return {"results": search_in_workspace(conv_id, q)}


@router.post("/conversations/{conv_id}/assistant-message")
async def add_assistant_message(conv_id: str, request: Request):
    """桌面文件投送完成后由桌面客户端调用：往对话写一条助手消息（自动同步 Supabase，App 实时可见）。

    body: {"content": "…", "files": [可选附件元数据]}。
    """
    require_conv_access(request, conv_id)
    try:
        body = await request.json()
    except Exception:
        body = {}
    content = (str(body.get("content") or "")).strip()
    if not content:
        raise HTTPException(400, "content 不能为空")
    files = body.get("files") if isinstance(body.get("files"), list) else None
    tool_calls = _external_tool_trace(body.get("tool_calls")) or None
    sources, groundings = _external_evidence_contract(content, body.get("sources"))
    client_message_id = str(body.get("client_message_id") or "").strip()
    if client_message_id and not re.fullmatch(r"[A-Za-z0-9._:-]{1,128}", client_message_id):
        raise HTTPException(422, "invalid_client_message_id")
    try:
      mid = db.create_message(conv_id, "assistant", content, files=files,
                            tool_calls=tool_calls, sources=sources or None,
                            groundings=groundings or None,
                            message_id=client_message_id or None)
    except ValueError:
      raise HTTPException(409, "client_message_id_conflict")
    return {"ok": True, "id": mid}


@router.post("/conversations/{conv_id}/user-message")
async def add_user_message(conv_id: str, request: Request):
    """往对话追加一条用户消息，但**不触发 Agent**（V209 澄清气泡）。

    用途：语音澄清——用户的原话与澄清问题都作为真实消息持久化（跨端同步、
    历史完整），Agent 只在用户补完信息后的下一轮才启动。body: {"content": "…"}。
    """
    require_conv_access(request, conv_id)
    try:
        body = await request.json()
    except Exception:
        body = {}
    content = (str(body.get("content") or "")).strip()
    if not content:
        raise HTTPException(400, "content 不能为空")
    client_message_id = str(body.get("client_message_id") or "").strip()
    if client_message_id and not re.fullmatch(r"[A-Za-z0-9._:-]{1,128}", client_message_id):
        raise HTTPException(422, "invalid_client_message_id")
    try:
      mid = db.create_message(
          conv_id, "user", content, message_id=client_message_id or None,
      )
    except ValueError:
      raise HTTPException(409, "client_message_id_conflict")
    return {"ok": True, "id": mid}


@router.post("/conversations/migrate")
async def migrate_conversations(request: Request):
    """把旧身份名下的对话过户到当前登录账号（Supabase 身份）。

    场景：以前匿名/本地账号建的对话，换成 Supabase 账号登录后看不到了——调用本接口一键认领。
    body 可选 {"from": "anonymous" 或 "<旧uid>"}，默认 anonymous。
    """
    user = get_current_user(request)
    if not user or not user.get("uid"):
        raise HTTPException(401, "请先登录账号再迁移")
    to_uid = user["uid"]
    try:
        body = await request.json()
    except Exception:
        body = {}
    from_uid = (str(body.get("from") or "anonymous")).strip()
    if from_uid == to_uid:
        return {"ok": True, "moved": 0, "from": from_uid, "to": to_uid, "note": "源与目标相同"}
    moved = db.reassign_conversations(from_uid, to_uid)
    return {"ok": True, "moved": moved, "from": from_uid, "to": to_uid}


@router.get("/conversations/migrate/preview")
async def migrate_conversations_preview(request: Request):
    """预览：当前账号有多少对话、anonymous 旧身份下还有多少可迁移。"""
    user = get_current_user(request)
    if not user or not user.get("uid"):
        raise HTTPException(401, "请先登录账号")
    to_uid = user["uid"]
    return {
        "current_uid": to_uid,
        "mine": db.count_conversations(to_uid),
        "anonymous": db.count_conversations("anonymous"),
    }


def _office_rich(fpath, ext: str) -> dict | None:
    """把 office/pdf 渲染成可在右栏直接预览的结构（docx→HTML、xlsx→表、pptx→分页文本、pdf→交给前端 iframe）。
    任一步失败返回 None，前端回退到「下载」。仅依赖后端本就装好的 python-docx/openpyxl/python-pptx。"""
    try:
        p = str(fpath)
        if ext == ".docx":
            from hashmm.api.doc_preview import docx_to_html
            return {"kind": "html", "html": docx_to_html(p)}
        if ext == ".xlsx":
            from hashmm.api.doc_preview import xlsx_to_json
            return {"kind": "sheets", "sheets": (xlsx_to_json(p) or {}).get("sheets", [])}
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(p)
            slides = []
            for slide in prs.slides:
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        t = shape.text_frame.text.strip()
                        if t:
                            texts.append(t)
                slides.append("\n".join(texts))
            return {"kind": "slides", "slides": slides, "pages": len(slides)}
        if ext == ".pdf":
            return {"kind": "pdf"}
    except Exception:
        return None
    return None


@router.get("/conversations/{conv_id}/files/{filename}/preview")
async def preview_conv_file(conv_id: str, filename: str, request: Request):
    """Preview a file in conversation workspace (text + office/pdf rich preview)."""
    require_conv_access(request, conv_id)
    import urllib.parse as _up, os.path as _osp
    # 防御：有的客户端会把 ?token=... 拼进文件名段（→ 文件名超长 OSError）。一律剥掉查询/锚点+目录。
    filename = _osp.basename(_up.unquote(filename).split("?")[0].split("#")[0])
    fdir = db.conv_files_dir(conv_id)
    fpath = fdir / filename
    if not fpath.exists():
        raise HTTPException(404, "文件不存在")

    ext = fpath.suffix.lower()
    stat = fpath.stat()
    size = stat.st_size
    # V361：文件预览使用弱 ETag 做增量校验。客户端可先显示持久缓存，后台只带
    # If-None-Match 询问版本；未变化直接 304，不再重复解析/传输 Office 文档。
    # mtime_ns + size 对本项目所有写路径（原子替换/新生成）都是稳定版本标识。
    etag = f'W/"{stat.st_mtime_ns:x}-{size:x}"'
    cache_headers = {
        "ETag": etag,
        "Cache-Control": "private, no-cache",
        "X-HashMM-File-Version": f"{stat.st_mtime_ns}:{size}",
    }
    if request.headers.get("if-none-match", "").strip() == etag:
        return Response(status_code=304, headers=cache_headers)
    binary_exts = {".pptx", ".docx", ".xlsx", ".pdf", ".zip", ".gz", ".tar",
                   ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".ico"}

    if ext in binary_exts or size > 1000000:
        # office/pdf 给出富预览结构；其余二进制仍只给下载。
        rich = _office_rich(fpath, ext) if ext in {".docx", ".xlsx", ".pptx", ".pdf"} else None
        payload = {
            "filename": filename, "ext": ext[1:], "size": size,
            "binary": True, "content": None, "lines": None,
            "rich": rich,
            "download_url": f"/api/conversations/{conv_id}/download/{filename}"
        }
        return JSONResponse(payload, headers=cache_headers)

    try:
        content = fpath.read_text(encoding="utf-8", errors="replace")
        lang_map = {".py": "python", ".js": "javascript", ".ts": "typescript",
                    ".java": "java", ".cpp": "cpp", ".c": "c", ".go": "go",
                    ".rs": "rust", ".html": "html", ".css": "css", ".json": "json",
                    ".md": "markdown", ".sh": "bash", ".yml": "yaml", ".yaml": "yaml"}
        payload = {
            "filename": filename, "ext": ext[1:], "size": size,
            "language": lang_map.get(ext, "text"),
            "content": content[:100000], "lines": len(content.splitlines()),
            "binary": False,
            "sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "download_url": f"/api/conversations/{conv_id}/download/{filename}"
        }
        return JSONResponse(payload, headers=cache_headers)
    except Exception as e:
        return JSONResponse({"filename": filename, "ext": ext[1:], "size": size,
                             "binary": True, "content": None, "error": str(e),
                             "download_url": f"/api/conversations/{conv_id}/download/{filename}"},
                            headers=cache_headers)


# ── V217: 画布版本历史落盘 ────────────────────────────────────────
# 走会话文件同一条链路：同 require_conv_access 鉴权、同 conv_files_dir 目录、
# 同 basename 防穿越。存放在工作区隐藏侧车 .wc-versions/<文件名>.json（list_conv_files
# 只列 is_file() 的顶层项，目录天然不进列表，不污染用户文件视图）。
_WC_VER_CAP = 10                    # 每画布最多保留版本数（与前端一致）
_WC_VER_HTML_CAP = 2_000_000        # 单版本 HTML 上限 2MB（与 preview 文本上限同量级）


def _wc_ver_path(conv_id: str, filename: str) -> Path:
    import urllib.parse as _up, os.path as _osp
    fname = _osp.basename(_up.unquote(filename).split("?")[0].split("#")[0])
    d = db.conv_files_dir(conv_id) / ".wc-versions"
    d.mkdir(parents=True, exist_ok=True)
    return d / (fname + ".json")


# ── V222: 会话文件写入（画布 2.0 地基）──────────────────────────
# 用途：①「画布」点击即开——前端用模板直接落一个 .html 进工作区；
#      ②画布就地编辑——用户改完 wc:save 回写原文件（版本历史另有 canvas-versions 快照链）。
# 防线与读取端完全同套：require_conv_access + basename 防穿越 + 大小上限 + 原子替换写。
_WRITE_EXT_ALLOW = {".html", ".htm", ".svg", ".md", ".txt", ".json", ".csv"}
_WRITE_SIZE_CAP = 2_000_000
_CANVAS_PATCH_OP_CAP = 24
_CANVAS_PATCH_HISTORY_CAP = 200


def _content_sha256(content: str) -> str:
    return hashlib.sha256(content.encode("utf-8")).hexdigest()


def _canvas_patch_path(conv_id: str, filename: str) -> Path:
    import os.path as _osp
    from urllib.parse import unquote as _unq
    fname = _osp.basename(_unq(filename).split("?")[0].split("#")[0])
    directory = db.conv_files_dir(conv_id) / ".canvas-patches"
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{fname}.json"


def _append_canvas_patch_audit(conv_id: str, filename: str, entry: dict) -> None:
    path = _canvas_patch_path(conv_id, filename)
    items: list[dict] = []
    try:
        if path.exists():
            loaded = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(loaded, list):
                items = [item for item in loaded if isinstance(item, dict)]
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        log_suppressed(_obs_logger, exc)
    items.append(entry)
    items = items[-_CANVAS_PATCH_HISTORY_CAP:]
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(items, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    os.replace(tmp, path)


@router.put("/conversations/{conv_id}/files/{filename}")
async def put_conv_file(conv_id: str, filename: str, request: Request):
    """写入/覆盖一个文本类会话文件。二进制（pptx/pdf 等）不走此口（画布 2.0 路线图 P2）。"""
    require_conv_access(request, conv_id)
    import os.path as _osp
    from urllib.parse import unquote as _unq
    fname = _osp.basename(_unq(filename).split("?")[0].split("#")[0])
    ext = _osp.splitext(fname)[1].lower()
    if ext not in _WRITE_EXT_ALLOW:
        raise HTTPException(415, f"该扩展名不允许写入（允许: {', '.join(sorted(_WRITE_EXT_ALLOW))}）")
    try:
        body = await request.json()
    except Exception:
        raise HTTPException(400, "需要 JSON 请求体 {content}")
    content = str((body or {}).get("content") or "")
    if not content:
        raise HTTPException(400, "content 不能为空")
    if len(content.encode("utf-8", "ignore")) > _WRITE_SIZE_CAP:
        raise HTTPException(413, "内容超过 2MB 上限")
    fdir = db.conv_files_dir(conv_id)
    fdir.mkdir(parents=True, exist_ok=True)
    fp = fdir / fname
    current = fp.read_text(encoding="utf-8", errors="replace") if fp.exists() else ""
    current_sha256 = _content_sha256(current) if fp.exists() else ""
    base_sha256 = str((body or {}).get("base_sha256") or "").strip().lower()
    lock_session = str((body or {}).get("lock_session") or "").strip()
    # Creating a new artifact remains compatible with the existing Chat path.
    # Overwriting an existing canvas from the visual editor is a privileged
    # commit: the writer must own the live lease and the exact base revision.
    if fp.exists() and (base_sha256 or lock_session):
        from hashmm.api.routes.canvas_locks import require_canvas_lock
        require_canvas_lock(conv_id, fname, lock_session)
        if not re.fullmatch(r"[0-9a-f]{64}", base_sha256):
            raise HTTPException(400, "base_sha256 必须是 64 位十六进制摘要")
        if base_sha256 != current_sha256:
            raise HTTPException(409, {
                "code": "canvas_revision_conflict",
                "message": "画布已在其他位置更新，请刷新并合并后再保存",
                "current_sha256": current_sha256,
            })
    tmp = fp.with_suffix(fp.suffix + ".tmp")
    tmp.write_text(content, encoding="utf-8")
    os.replace(tmp, fp)
    new_sha256 = _content_sha256(content)
    try:
        db.add_conversation_file(conv_id, fname, size=len(content.encode("utf-8", "ignore")), mime_type="text/html" if ext in (".html", ".htm") else "text/plain")
    except Exception as e:
        log_suppressed(_obs_logger, e)
    return {"ok": True, "filename": fname,
            "download_url": f"/api/conversations/{conv_id}/download/{fname}",
            "size": len(content), "sha256": new_sha256}


@router.patch("/conversations/{conv_id}/files/{filename}/canvas-blocks")
async def patch_canvas_blocks(conv_id: str, filename: str, request: Request):
    """Apply deterministic, revision-checked block patches to an HTML canvas.

    The model or UI must provide the exact old fragment for every replacement.
    A missing or non-unique anchor fails the complete transaction; no partial
    mutation is committed.  Audit history stores hashes/lengths only, never the
    potentially sensitive page content.
    """
    conv = require_conv_access(request, conv_id)
    import os.path as _osp
    from urllib.parse import unquote as _unq
    fname = _osp.basename(_unq(filename).split("?")[0].split("#")[0])
    if _osp.splitext(fname)[1].lower() not in {".html", ".htm"}:
        raise HTTPException(415, "块级补丁仅支持 HTML 画布")
    try:
        body = await request.json()
    except Exception:
        raise HTTPException(400, "需要 JSON 请求体")
    base_sha256 = str((body or {}).get("base_sha256") or "").strip().lower()
    lock_session = str((body or {}).get("lock_session") or "").strip()
    idempotency_key = _bounded_text((body or {}).get("idempotency_key"), 96, one_line=True)
    operations = (body or {}).get("operations")
    if not re.fullmatch(r"[0-9a-f]{64}", base_sha256):
        raise HTTPException(400, "base_sha256 必须是 64 位十六进制摘要")
    if not isinstance(operations, list) or not operations or len(operations) > _CANVAS_PATCH_OP_CAP:
        raise HTTPException(400, f"operations 必须包含 1-{_CANVAS_PATCH_OP_CAP} 个补丁")
    from hashmm.api.routes.canvas_locks import require_canvas_lock
    require_canvas_lock(conv_id, fname, lock_session)
    fp = db.conv_files_dir(conv_id) / fname
    if not fp.exists() or not fp.is_file():
        raise HTTPException(404, "文件不存在")
    content = fp.read_text(encoding="utf-8", errors="replace")
    current_sha256 = _content_sha256(content)
    if current_sha256 != base_sha256:
        raise HTTPException(409, {
            "code": "canvas_revision_conflict",
            "message": "画布基线版本已变化，补丁未应用",
            "current_sha256": current_sha256,
        })
    audit_ops: list[dict] = []
    next_content = content
    for index, raw in enumerate(operations):
        if not isinstance(raw, dict):
            raise HTTPException(400, f"operations[{index}] 必须是对象")
        op = str(raw.get("op") or "replace").strip()
        block_id = _bounded_text(raw.get("block_id"), 128, one_line=True) or f"block-{index + 1}"
        old = str(raw.get("old_html") or "")
        new = str(raw.get("new_html") or "")
        if op not in {"replace", "insert_after"}:
            raise HTTPException(400, f"operations[{index}].op 不支持")
        if not old or len(old.encode("utf-8")) > 500_000 or len(new.encode("utf-8")) > 500_000:
            raise HTTPException(400, f"operations[{index}] 片段为空或过大")
        occurrences = next_content.count(old)
        if occurrences != 1:
            raise HTTPException(409, {
                "code": "canvas_patch_anchor_conflict",
                "message": f"块 {block_id} 的锚点匹配 {occurrences} 次，整批补丁未应用",
                "block_id": block_id,
            })
        replacement = new if op == "replace" else old + new
        next_content = next_content.replace(old, replacement, 1)
        audit_ops.append({
            "op": op,
            "block_id": block_id,
            "old_sha256": _content_sha256(old),
            "new_sha256": _content_sha256(new),
            "old_bytes": len(old.encode("utf-8")),
            "new_bytes": len(new.encode("utf-8")),
        })
    if len(next_content.encode("utf-8")) > _WRITE_SIZE_CAP:
        raise HTTPException(413, "补丁后的画布超过 2MB 上限")
    new_sha256 = _content_sha256(next_content)
    if new_sha256 == current_sha256:
        return {"ok": True, "dedup": True, "sha256": new_sha256, "applied": 0}
    tmp = fp.with_suffix(fp.suffix + ".patch.tmp")
    tmp.write_text(next_content, encoding="utf-8")
    os.replace(tmp, fp)
    actor = get_current_user(request) or {}
    _append_canvas_patch_audit(conv_id, fname, {
        "schema": "hashmm.canvas-patch-receipt.v1",
        "patch_id": str(uuid.uuid4()),
        "idempotency_key": idempotency_key,
        "base_sha256": current_sha256,
        "result_sha256": new_sha256,
        "actor_id": str(actor.get("uid") or conv.get("user_id") or "anonymous")[:128],
        "created_at": time.time(),
        "operations": audit_ops,
    })
    try:
        db.add_conversation_file(
            conv_id, fname, size=len(next_content.encode("utf-8")), mime_type="text/html"
        )
    except Exception as exc:
        log_suppressed(_obs_logger, exc)
    return {
        "ok": True, "dedup": False, "sha256": new_sha256,
        "base_sha256": current_sha256, "applied": len(audit_ops),
    }


@router.get("/conversations/{conv_id}/files/{filename}/canvas-versions")
async def get_canvas_versions(conv_id: str, filename: str, request: Request):
    """读取一张画布的版本历史（跨会话持久）。文件不存在/损坏一律回空列表，绝不拦主流程。"""
    require_conv_access(request, conv_id)
    p = _wc_ver_path(conv_id, filename)
    try:
        if p.exists():
            data = json.loads(p.read_text(encoding="utf-8"))
            if isinstance(data, list):
                return {"versions": data[-_WC_VER_CAP:]}
    except Exception as e:
        log_suppressed(_obs_logger, e)
    return {"versions": []}


@router.put("/conversations/{conv_id}/files/{filename}/canvas-versions")
async def put_canvas_version(conv_id: str, filename: str, request: Request):
    """追加一版画布快照。幂等：与最后一版相同的 html 不重复入史；超 10 版滚动淘汰最旧。"""
    require_conv_access(request, conv_id)
    try:
        body = await request.json()
    except Exception:
        raise HTTPException(400, "需要 JSON 请求体")
    html = str((body or {}).get("html") or "")
    if not html:
        raise HTTPException(400, "html 不能为空")
    if len(html) > _WC_VER_HTML_CAP:
        raise HTTPException(413, "单版本超过 2MB 上限")
    p = _wc_ver_path(conv_id, filename)
    versions: list = []
    try:
        if p.exists():
            loaded = json.loads(p.read_text(encoding="utf-8"))
            if isinstance(loaded, list):
                versions = loaded
    except Exception as e:
        log_suppressed(_obs_logger, e)
    if versions and isinstance(versions[-1], dict) and versions[-1].get("html") == html:
        return {"versions": len(versions), "dedup": True}
    versions.append({"ts": int(time.time() * 1000), "html": html})
    versions = versions[-_WC_VER_CAP:]
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(versions, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, p)
    return {"versions": len(versions)}


@router.get("/conversations/{conv_id}/files/{filename}/view", response_class=HTMLResponse)
async def view_conv_file(conv_id: str, filename: str, request: Request):
    """返回适配手机的完整 HTML 预览页：App 内置 WebView 直接打开、不跳浏览器。
    docx→HTML、xlsx→表格、pptx→分页、txt/code→等宽文本、pdf→pdf.js 渲染、图片→<img>。"""
    require_conv_access(request, conv_id)
    import urllib.parse, html as _html, os.path as _osp
    fname = _osp.basename(urllib.parse.unquote(filename).split("?")[0].split("#")[0])
    fdir = db.conv_files_dir(conv_id)
    fpath = fdir / fname
    if not fpath.exists():
        return HTMLResponse("<h3 style='font-family:sans-serif;padding:24px'>文件不存在</h3>", status_code=404)
    ext = fpath.suffix.lower()
    tok = request.query_params.get("token", "")
    dl = f"/api/conversations/{conv_id}/download/{urllib.parse.quote(fname)}"
    dl_tok = dl + (("?token=" + urllib.parse.quote(tok)) if tok else "")
    head = ("<!doctype html><html><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width,initial-scale=1,maximum-scale=5'>"
            "<style>body{margin:0;padding:16px;font-family:system-ui,-apple-system,sans-serif;"
            "color:#1a1a1a;background:#fff;line-height:1.7;font-size:15px;-webkit-text-size-adjust:100%}"
            "table{border-collapse:collapse;width:100%;font-size:13px;margin:8px 0;display:block;overflow-x:auto}"
            "td,th{border:1px solid #e5e7eb;padding:6px 10px;text-align:left;white-space:nowrap}"
            "th{background:#f9fafb;font-weight:600}"
            "pre{white-space:pre-wrap;word-break:break-word;background:#f6f8fa;padding:12px;border-radius:8px;font-size:13px}"
            ".slide{border:1px solid #e5e7eb;border-radius:10px;padding:14px;margin:10px 0}"
            ".slide .n{font-size:11px;color:#888;margin-bottom:6px}img{max-width:100%}"
            "canvas{max-width:100%;margin-bottom:8px;box-shadow:0 1px 4px rgba(0,0,0,.12)}"
            "h1,h2,h3{margin:.6em 0 .3em}a{color:#2563eb}</style></head><body>")
    tail = "</body></html>"
    try:
        if ext == ".docx":
            from hashmm.api.doc_preview import docx_to_html
            body = docx_to_html(str(fpath))
        elif ext == ".xlsx":
            from hashmm.api.doc_preview import xlsx_to_json
            sheets = (xlsx_to_json(str(fpath)) or {}).get("sheets", [])
            parts = []
            for sh in sheets:
                parts.append(f"<h3>{_html.escape(str(sh.get('name', '')))}</h3><table>")
                headers = sh.get("headers") or []
                if headers:
                    parts.append("<tr>" + "".join(f"<th>{_html.escape(str(h))}</th>" for h in headers) + "</tr>")
                for row in (sh.get("rows") or []):
                    parts.append("<tr>" + "".join(f"<td>{_html.escape('' if c is None else str(c))}</td>" for c in row) + "</tr>")
                parts.append("</table>")
            body = "".join(parts) or "<p>空表格</p>"
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(fpath))
            parts = []
            for i, slide in enumerate(prs.slides):
                txt = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        t = shape.text_frame.text.strip()
                        if t:
                            txt.append(_html.escape(t))
                parts.append(f"<div class='slide'><div class='n'>第 {i + 1} 页</div>" + "<br>".join(txt) + "</div>")
            body = "".join(parts) or "<p>空演示文稿</p>"
        elif ext == ".pdf":
            body = (f"<div id='pdf'></div><p id='fb' style='display:none'>无法在线渲染，<a href='{dl_tok}'>点此下载查看</a></p>"
                    "<script src='https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/pdf.min.js'></script>"
                    "<script>try{pdfjsLib.GlobalWorkerOptions.workerSrc='https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/pdf.worker.min.js';"
                    f"pdfjsLib.getDocument('{dl_tok}').promise.then(function(pdf){{"
                    "for(var i=1;i<=pdf.numPages;i++){(function(n){pdf.getPage(n).then(function(p){"
                    "var vp=p.getViewport({scale:1.5});var c=document.createElement('canvas');var ctx=c.getContext('2d');"
                    "c.width=vp.width;c.height=vp.height;document.getElementById('pdf').appendChild(c);"
                    "p.render({canvasContext:ctx,viewport:vp});});})(i);}"
                    "}).catch(function(){document.getElementById('fb').style.display='block';});}"
                    "catch(e){document.getElementById('fb').style.display='block';}</script>")
        elif ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"):
            body = f"<img src='{dl_tok}'>"
        else:
            content = fpath.read_text(encoding="utf-8", errors="replace")[:200000]
            body = f"<pre>{_html.escape(content)}</pre>"
    except Exception as e:
        body = f"<p>预览失败：{_html.escape(str(e)[:120])}</p><p><a href='{dl_tok}'>点此下载查看</a></p>"
    return HTMLResponse(head + body + tail)




@router.get("/conversations/{conv_id}/export/jupyter")
async def export_jupyter(conv_id: str, request: Request):
    """Export conversation as Jupyter Notebook."""
    require_conv_access(request, conv_id)
    nb = db.export_as_jupyter(conv_id)
    import tempfile
    with tempfile.NamedTemporaryFile(mode='w', suffix='.ipynb', delete=False, encoding='utf-8') as f:
        json.dump(nb, f, ensure_ascii=False)
        tmp = f.name
    return FileResponse(tmp, filename=f"{conv_id[:8]}_notebook.ipynb",
                       media_type="application/x-ipynb+json")


@router.get("/conversations/{conv_id}/export/markdown")
async def export_markdown(conv_id: str, request: Request):
    """Export conversation as Markdown."""
    require_conv_access(request, conv_id)
    md = db.export_as_markdown(conv_id)
    return Response(content=md, media_type="text/markdown",
                   headers={"Content-Disposition": f"attachment; filename={conv_id[:8]}.md"})


@router.get("/conversations/{conv_id}/export/html")
async def export_html(conv_id: str, request: Request):
    """Export conversation as standalone HTML."""
    require_conv_access(request, conv_id)
    html = db.export_as_html(conv_id)
    return Response(content=html, media_type="text/html",
                   headers={"Content-Disposition": f"attachment; filename={conv_id[:8]}.html"})




@router.get("/conversations/{conv_id}/dependencies")
async def analyze_deps(conv_id: str, request: Request):
    """Analyze project dependencies in workspace."""
    require_conv_access(request, conv_id)
    from hashmm.api.workspace import analyze_dependencies
    return analyze_dependencies(conv_id)



@router.post("/conversations/{conv_id}/share")
async def share_conversation(conv_id: str, request: Request):
    """Create a share link for a conversation."""
    require_conv_access(request, conv_id)
    user = get_current_user(request)
    user_id = user["uid"] if user else "anonymous"
    body = await request.json() if request.headers.get("content-type") == "application/json" else {}
    code = db.create_share(conv_id, user_id,
                          password=body.get("password"),
                          expires_hours=body.get("expires_hours", 72))
    return {"share_code": code, "url": f"/shared/{code}"}


@router.post("/conversations/{conv_id}/execute")
async def execute_code_in_conv(conv_id: str, request: Request):
    """Execute code within conversation workspace."""
    require_conv_access(request, conv_id)
    body = await request.json()
    code_text = body.get("code", "")
    if not code_text:
        return {"error": "代码为空"}
    user_id = resolve_user_id(request)
    args = {"code": code_text, "timeout": 15}
    from hashmm.agent.execution_scope import build_root_scope
    execution_scope = build_root_scope(
        owner_id=user_id,
        conversation_id=conv_id,
        run_id="direct-exec-" + uuid.uuid4().hex[:16],
        allowed_tools=["execute_code"],
        approval_mode="read_only",
        network_mode="deny",
        allow_subagents=False,
        max_tool_calls=1,
        max_workers=0,
    )
    result = execute_tool("execute_code", args, {
        "user_id": user_id,
        "conv_id": conv_id,
        "session_id": conv_id,
        "cwd": str(db.conv_files_dir(conv_id)),
        "execution_scope": execution_scope,
        # This endpoint is invoked by an explicit user click after the object
        # owner check above.  It approves this exact call, never future calls.
        "approved": True,
    })
    from hashmm.api.tool_result import parse_tool_result
    normalized = parse_tool_result(result)
    return {
        "ok": normalized.success,
        "status": "done" if normalized.success else "error",
        "output": normalized.content if normalized.success else "",
        "error": normalized.error if not normalized.success else None,
        "files": normalized.files,
        "metrics": normalized.metrics,
    }




@router.patch("/conversations/{conv_id}/messages/{msg_id}/edit")
async def edit_message(conv_id: str, msg_id: str, request: Request):
    """Edit a user message (deletes all subsequent messages)."""
    require_conv_access(request, conv_id)
    body = await request.json()
    new_content = body.get("content", "")
    if not new_content:
        raise HTTPException(400, "内容不能为空")
    
    # Update the message
    db.update_message(msg_id, content=new_content)
    
    # Delete all messages after this one
    with db._conn() as c:
        msg = c.execute("SELECT created_at FROM messages WHERE id=?", (msg_id,)).fetchone()
        if msg:
            c.execute("DELETE FROM messages WHERE conv_id=? AND created_at > ?",
                      (conv_id, msg["created_at"]))
            c.execute("DELETE FROM conversation_compactions WHERE conv_id=?", (conv_id,))
    
    return {"ok": True, "message_id": msg_id}


@router.post("/conversations/{conv_id}/messages/{msg_id}/regenerate")
async def regenerate_message(conv_id: str, msg_id: str, request: Request):
    """Regenerate response from a specific message point."""
    require_conv_access(request, conv_id)
    with db._conn() as c:
        # Find the user message before this assistant message
        msg = c.execute("SELECT * FROM messages WHERE id=?", (msg_id,)).fetchone()
        if not msg:
            raise HTTPException(404, "消息不存在")
        
        # Delete this message and all after it
        c.execute("DELETE FROM messages WHERE conv_id=? AND created_at >= ?",
                  (conv_id, msg["created_at"]))
        c.execute("DELETE FROM conversation_compactions WHERE conv_id=?", (conv_id,))
    
    return {"ok": True, "conv_id": conv_id}




@router.post("/conversations/{cid}/assign-project")
async def assign_project(cid: str, request: Request):
    require_conv_access(request, cid)
    user = require_auth(request)
    body = await request.json()
    if not db.assign_conv_to_project(
        cid,
        body.get("project_id"),
        str(user.get("uid") or ""),
    ):
        # A foreign/missing project and conversation share the same bounded
        # response so object IDs cannot be enumerated across owners.
        raise HTTPException(404, "对话或项目不存在")
    return {"ok": True}




@router.post("/conversations/{conv_id}/tags")
async def add_conv_tag(conv_id: str, request: Request):
    require_conv_access(request, conv_id)
    body = await request.json()
    tag = body.get("tag", "").strip()
    if not tag:
        raise HTTPException(400, "标签不能为空")
    db.add_tag(conv_id, tag)
    return {"ok": True, "tags": db.get_tags(conv_id)}


@router.delete("/conversations/{conv_id}/tags/{tag}")
async def remove_conv_tag(conv_id: str, tag: str, request: Request):
    require_conv_access(request, conv_id)
    db.remove_tag(conv_id, tag)
    return {"ok": True}


@router.post("/conversations/{conv_id}/fork/{msg_id}")
async def fork_conversation(conv_id: str, msg_id: str, request: Request):
    """Fork conversation from a specific message. Creates a new conversation
    with all messages up to and including msg_id, owned by the caller."""
    require_conv_access(request, conv_id)
    uid = resolve_user_id(request)

    # v17 Phase 67: the old body called db.create_conversation(uid, title)
    # (argument order is (conv_id, user_id, title) → it built a broken row keyed
    # by the user id) and db.add_message (which doesn't exist). Delegate to the
    # correct, already-tested helper instead.
    with db._conn() as c:
        exists = c.execute("SELECT 1 FROM messages WHERE id=? AND conv_id=?",
                           (msg_id, conv_id)).fetchone()
    if not exists:
        raise HTTPException(404, "消息不存在")

    new_id = db.fork_conversation(conv_id, msg_id, uid)
    if not new_id:
        raise HTTPException(404, "对话不存在")
    copied = len(db.get_messages(new_id, limit=500))
    return {"ok": True, "new_conv_id": new_id, "messages_copied": copied}


@router.post("/conversations/{conv_id}/files/create")
async def create_file_in_conv(conv_id: str, request: Request):
    """Create a file in conversation workspace (from CodeBlock button)."""
    require_conv_access(request, conv_id)
    body = await request.json()
    filename = body.get("filename", "untitled.py")
    content = body.get("content", "")
    if not content:
        raise HTTPException(400, "内容为空")
    result = execute_tool("create_file", {"filename": filename, "content": content}, {"session_id": conv_id})
    return {"ok": "OK" in result, "result": result, "download_url": f"/api/conversations/{conv_id}/download/{filename}"}


async def _write_conversation_upload(fpath: Path, content: bytes) -> dict | None:
    """Write an upload without replacing a valid Office file by invalid bytes."""
    ext = fpath.suffix.lower()
    if ext not in {".docx", ".pptx", ".xlsx"}:
        fpath.write_bytes(content)
        return None
    staged = fpath.with_name(f".{fpath.stem}.upload-{uuid.uuid4().hex}{ext}")
    try:
        staged.write_bytes(content)
        from hashmm.agent.office_artifacts import inspect_office
        report = await run_in_threadpool(inspect_office, staged)
        if not report.get("verified"):
            raise HTTPException(422, "Office 文件无法完整解析，未覆盖当前任务中的原文件")
        os.replace(staged, fpath)
        report["file"]["name"] = fpath.name
        return report
    finally:
        try:
            if staged.exists():
                staged.unlink()
        except OSError:
            pass


@router.post("/conversations/{conv_id}/upload")
async def upload_file_to_conv(conv_id: str, request: Request, file: UploadFile = File(...)):
    """Upload a validated file to the caller-owned conversation workspace."""
    conversation = require_conv_access(request, conv_id)
    raw_name = file.filename or "uploaded_file"
    safe_name = Path(raw_name).name
    if not safe_name or safe_name != raw_name or safe_name in (".", ".."):
        raise HTTPException(400, "非法文件名")
    content = await file.read()
    from hashmm.api.core.safety import validate_upload
    valid, reason = validate_upload(safe_name, len(content))
    if not valid:
        raise HTTPException(400, reason)
    fdir = db.conv_files_dir(conv_id)
    fdir.mkdir(parents=True, exist_ok=True)
    fpath = fdir / safe_name
    content_sha256 = hashlib.sha256(content).hexdigest()
    if fpath.is_file():
        try:
            existing_sha256 = hashlib.sha256(fpath.read_bytes()).hexdigest()
        except OSError:
            existing_sha256 = ""
        if existing_sha256 != content_sha256:
            suffix = fpath.suffix
            stem = fpath.stem[: max(1, 220 - len(suffix))]
            fpath = fdir / f"{stem}-{content_sha256[:10]}{suffix}"
    office_artifact = await _write_conversation_upload(fpath, content)
    db.add_conversation_file(conv_id, fpath.name, len(content), file.content_type or "")
    # Storage success and model readability are separate facts. Parse the
    # owner-scoped committed bytes now (content-addressed cache makes the Chat
    # turn cheap) and return an observable state to every client.
    try:
        parsed_resource = await run_in_threadpool(
            parse_resource, fpath, display_name=fpath.name
        )
        resource = resource_summary(parsed_resource)
    except Exception as exc:
        resource = {
            "filename": fpath.name,
            "sha256": hashlib.sha256(content).hexdigest(),
            "parse_state": "error",
            "page_count": 0,
            "readable_pages": 0,
            "readable_ratio": 0.0,
            "warnings": [f"解析失败：{type(exc).__name__}: {exc}"],
        }
    
    ocr_job = None
    if resource.get("parse_state") == "needs_ocr":
        try:
            from hashmm.pipeline import ocr_queue
            ocr_job = await run_in_threadpool(
                lambda: ocr_queue.enqueue(
                    user_id=str(conversation.get("user_id") or ""),
                    conv_id=conv_id,
                    project_id=str(conversation.get("project_id") or ""),
                    filename=fpath.name,
                    source_path=str(fpath),
                    sha256=str(resource.get("sha256") or ""),
                    resource_id=str(resource.get("resource_id") or ""),
                    resource_revision=int(resource.get("resource_revision") or 1),
                )
            )
        except Exception as exc:
            warnings = list(resource.get("warnings") or [])
            warnings.append(f"OCR queue admission failed: {type(exc).__name__}")
            resource["warnings"] = warnings

    # Check if it's an image for vision analysis
    ext = fpath.suffix.lower()
    is_image = ext in ('.png', '.jpg', '.jpeg', '.gif', '.webp')
    
    return {
        "ok": True,
        "filename": fpath.name,
        "size": len(content),
        # Content receipt for desktop/App incremental caches. This is the hash
        # of the exact validated bytes committed to the owned workspace, not a
        # model assertion or a client-provided checksum.
        "sha256": content_sha256,
        "is_image": is_image,
        "download_url": f"/api/conversations/{conv_id}/download/{fpath.name}",
        "office_artifact": office_artifact,
        "resource": resource,
        "ocr_job": ocr_job,
    }


@router.get("/phone-file-requests")
async def list_phone_file_requests(request: Request):
    """App 轮询：列出"电脑端向手机要照片"的待办请求（pending, target=phone）。"""
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    from hashmm.api import supabase_sync as _sbs
    rows = _sbs.list_phone_requests(user.get("uid"))
    out = [{"id": r.get("id"), "conv_id": r.get("conv_id"), "query": r.get("query", "")}
           for r in rows if isinstance(r, dict) and r.get("id") and r.get("conv_id")]
    return {"requests": out}


@router.post("/phone-file-requests/{req_id}/status")
async def set_phone_file_request_status(req_id: str, request: Request):
    """App 回写：将某条手机照片请求标记为 processing/done/denied/error。"""
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    try:
        body = await request.json()
    except Exception:
        body = {}
    status = (body.get("status") or "done")[:20]
    from hashmm.api import supabase_sync as _sbs
    _sbs.set_request_status(req_id, status)
    return {"ok": True}


def _dispatch_or_warn(conv_id: str, uid, query: str) -> dict:
    """把请求写入本地 file_requests，桌面端常驻轮询接走执行（手机与电脑共用同一后端，无需 Supabase key）。
    若 Supabase 也配置了，则额外推一份，兼容异地后端/多设备；两条都失败安全。"""
    import uuid as _uuid
    q = (query or "")[:500]
    ok_local = False
    req_id = str(_uuid.uuid4())
    try:
        db.create_file_request(req_id, uid, conv_id, q, target="desktop")
        ok_local = True
    except Exception:
        ok_local = False
    try:
        from hashmm.api import supabase_sync as _sbs
        if _sbs.enabled():
            _sbs.push_file_request(uid, conv_id, q, target="desktop")
    except Exception:
        pass
    if not ok_local:
        try:
            db.create_message(conv_id, "assistant", "这条请求没能进入下发队列，请稍后重试。")
        except Exception:
            pass
        return {"ok": False, "reason": "enqueue_failed", "dispatched": q}
    # The queue row is the execution hand-off; the work runtime is the shared
    # user-facing source of truth consumed by Chat, desktop and App.  Keep the
    # legacy prefix out of titles and never copy runner arguments/results into
    # this projection.
    try:
        from hashmm.agent import work_runtime as _work_runtime
        is_browser = q.startswith("[[AGENT")
        visible_task = re.sub(r"^\[\[[A-Z_!]+\]\]\s*", "", q).strip()
        work_run = _work_runtime.create_run(
            user_id=str(uid or ""),
            kind="browser" if is_browser else "computer",
            source_id=req_id,
            conv_id=conv_id,
            title=visible_task[:120] or ("浏览器任务" if is_browser else "电脑任务"),
            status="queued",
            snapshot={
                "request_id": req_id,
                "target": "desktop",
                "capability": "browser_use" if is_browser else (
                    "computer_use" if q.startswith("[[") else "file_transfer"
                ),
                "delivery": "cross_device",
            },
        )
        _work_runtime.append_event(
            work_run["id"], user_id=str(uid or ""), event_type="device_dispatch",
            status="queued", summary="已发送到桌面端执行队列",
            payload={"target": "desktop", "request_id": req_id},
        )
        work_run_id = work_run["id"]
    except Exception as exc:
        # Queue admission remains successful even if observability fails.  The
        # runner must not execute twice merely to repair its projection.
        log_suppressed(_obs_logger, exc, "work runtime admission")
        work_run_id = ""
    return {
        "ok": True, "request_id": req_id, "work_run_id": work_run_id,
        "status": "pending", "dispatched": q,
    }


@router.post("/conversations/{conv_id}/request-file")
async def request_file_from_desktop(conv_id: str, request: Request):
    """App「从电脑取文件」一键下发：写一条 file_requests(target=desktop)，桌面端常驻轮询器接走、
    在 桌面/下载/文档 里按文件名匹配并上传回本会话。零新增后端逻辑——复用桌面已有的投送链路。"""
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    try:
        body = await request.json()
    except Exception:
        body = {}
    query = str(body.get("query") or body.get("filename") or "").strip()
    if not query:
        raise HTTPException(status_code=400, detail="缺少要取的文件名/描述（query）")
    require_conv_access_or_create(request, conv_id, "从电脑取文件")
    db.create_message(conv_id, "user", f"【取电脑文件】{query}")
    return _dispatch_or_warn(conv_id, user.get("uid"), query)


@router.post("/conversations/{conv_id}/computer-task")
async def computer_task_from_app(conv_id: str, request: Request):
    """App「让电脑做事」：写一条 file_requests(target=desktop)，query 带 [[CU]] 前缀，
    桌面端轮询器识别前缀后执行（当前支持：列出 桌面/下载/文档 文件清单），结果回写本会话。
    复用取文件的同一条投送链路，零新增表结构。"""
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    try:
        body = await request.json()
    except Exception:
        body = {}
    task = str(body.get("task") or body.get("query") or "").strip()
    if not task:
        raise HTTPException(status_code=400, detail="缺少要让电脑做的事（task）")
    requested_kind = str(body.get("kind") or "computer_use").strip().lower()
    # App/desktop use the same public capability names.  Keep the old wire
    # prefixes for installed runners so upgrades remain backwards compatible.
    kind = {
        "browser_use": "agent",
        "computer_use": "cu",
    }.get(requested_kind, requested_kind)
    prefix = {
        "agent": "[[AGENT]] ", "cmd": "[[CMD]] ", "auto": "[[AUTO]] ",
        "exec": "[[EXEC]] ", "exec_ok": "[[EXEC!]] ", "agent_ok": "[[AGENT!]] ",
        "agent_resume": "[[AGENT_RESUME]] ", "mem": "[[MEM]] ",
        "mem_forget": "[[MEM_FORGET]] ", "mem_clear": "[[MEM_CLEAR]] ",
        "mem_pref_forget": "[[MEM_PREF_FORGET]] ", "mem_field_clear": "[[MEM_FIELD_CLEAR]] ",
        "mem_field_set": "[[MEM_FIELD_SET]] ",
        "seq": "[[SEQ]] ", "seq_resume": "[[SEQ_RESUME]] ", "task_cancel": "[[TASK_CANCEL]] ",
    }.get(kind, "[[CU]] ")
    require_conv_access_or_create(request, conv_id, "电脑任务")
    label = "【浏览器调研】" if kind == "agent" else "【电脑任务】"
    db.create_message(conv_id, "user", label + task)
    res = _dispatch_or_warn(conv_id, user.get("uid"), prefix + task)
    res["kind"] = requested_kind
    res["wire_kind"] = kind
    if res.get("ok"):
        res["dispatched"] = task[:480]
    return res


@router.get("/conversations/{conv_id}/computer-requests")
async def conversation_computer_requests(conv_id: str, request: Request, limit: int = 20):
    """Persisted desktop-task state for Chat and App task pages.

    Ownership is checked before querying. Missing and unauthorized
    conversations both return the same 404 through require_conv_access.
    """
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    require_conv_access(request, conv_id)
    rows = db.list_file_requests_for_conversation(user.get("uid"), conv_id, limit=limit)
    return {"requests": rows}


@router.get("/file-requests/pending")
async def list_pending_requests(request: Request):
    """桌面端常驻轮询：取本账号待处理的电脑任务/取文件请求（本地直达，无需 Supabase）。"""
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    target = (request.query_params.get("target") or "desktop").strip() or "desktop"
    try:
        rows = db.list_pending_file_requests(user.get("uid"), target=target, limit=5)
    except Exception:
        rows = []
    return {"requests": rows}


@router.patch("/file-requests/{req_id}")
async def patch_file_request(req_id: str, request: Request):
    """桌面端回写请求状态（processing / done / error），避免重复处理。"""
    user = get_current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="需要登录")
    try:
        body = await request.json()
    except Exception:
        body = {}
    status = str(body.get("status") or "").strip()
    if status not in ("pending", "processing", "done", "error"):
        raise HTTPException(status_code=400, detail="无效状态")
    try:
        updated = db.update_file_request(req_id, user.get("uid"), status)
    except Exception:
        updated = False
    if not updated:
        # Missing and another user's request deliberately share one response.
        raise HTTPException(status_code=404, detail="任务不存在")
    work_run_id = ""
    try:
        from hashmm.agent import work_runtime as _work_runtime
        owner = str(user.get("uid") or "")
        work_run = _work_runtime.get_run_by_source(req_id, owner)
        if work_run:
            work_run_id = work_run["id"]
            transition = {
                "pending": ("queued", "queued", "任务仍在等待桌面端领取"),
                "processing": ("started", "running", "桌面端已领取并开始执行"),
                "done": ("delivered", "delivered", "桌面端已完成并交付结果"),
                "error": ("failed", "failed", "桌面端执行失败，可从会话重试"),
            }[status]
            _work_runtime.append_event(
                work_run_id, user_id=owner, event_type=transition[0],
                status=transition[1], summary=transition[2],
                payload={"target": "desktop", "request_id": req_id},
                snapshot_updates={"desktop_request_status": status},
                # Desktop acknowledgements are retried after timeouts.  The
                # server status transition is a fact, not a counter: replaying
                # the same request/status must not create another event or
                # advance the cross-device revision.
                idempotency_key=f"desktop-request:{req_id}:{status}",
            )
    except Exception as exc:
        # State acknowledgement belongs to file_requests and has already been
        # committed.  A ledger projection failure is observable but must not
        # make the desktop retry the operation and duplicate its side effects.
        log_suppressed(_obs_logger, exc, "work runtime transition")
    response = {"ok": True}
    if work_run_id:
        response["work_run_id"] = work_run_id
    return response


@router.get("/privacy-mode")
async def get_privacy_mode(request: Request):
    """本地隐私模式状态。on=对话不上云; llm_local=默认模型是否本机; local_model_available=是否配了可切的本地模型。"""
    require_auth(request)
    from hashmm.api import privacy_mode as _pv
    return _pv.status()


@router.post("/privacy-mode")
async def set_privacy_mode(request: Request):
    require_auth(request)
    from hashmm.api import privacy_mode as _pv
    try:
        body = await request.json()
    except Exception:
        body = {}
    _pv.set_on(bool(body.get("on")))
    return {"ok": True, **_pv.status()}




@router.post("/conversations/{conv_id}/fork")
async def fork_conversation_api(conv_id: str, request: Request):
    """Fork conversation from a specific message."""
    require_conv_access(request, conv_id)
    uid = resolve_user_id(request)
    body = await request.json()
    msg_id = body.get("message_id", "")
    new_id = db.fork_conversation(conv_id, msg_id, uid)
    if not new_id:
        raise HTTPException(400, "分支失败")
    return {"ok": True, "conv_id": new_id, "url": f"/chat/{new_id}"}


@router.patch("/conversations/{conv_id}/prompt")
async def set_conv_custom_prompt(conv_id: str, request: Request):
    """Set per-conversation custom system prompt."""
    require_conv_access(request, conv_id)
    body = await request.json()
    prompt = body.get("prompt", "")
    with db._conn() as c:
        try:
            c.execute("ALTER TABLE conversations ADD COLUMN custom_prompt TEXT DEFAULT ''")
        except Exception: pass
        c.execute("UPDATE conversations SET custom_prompt=? WHERE id=?", (prompt, conv_id))
    return {"ok": True}


@router.get("/conversations/{conv_id}/prompt")
async def get_conv_custom_prompt(conv_id: str, request: Request):
    """Return the persisted owner-scoped conversation instruction."""
    conv = require_conv_access(request, conv_id)
    return {"prompt": str(conv.get("custom_prompt") or "")}



# ── v10.0: Message feedback (triggers skill creation) ──

class FeedbackRequest(BaseModel):
    feedback: str  # "up" or "down"
    message_content: str = ""
    query: str = ""


class FeedbackReviewRequest(BaseModel):
    rating: str  # up | down | clear
    reason_code: str = ""
    comment: str = ""


def _apply_feedback_learning(case: dict, uid: str) -> None:
    """Best-effort evolution hooks fed only with server-derived conversation data."""
    rating = str(case.get("rating") or "")
    query = str(case.get("query") or "")
    answer = str(case.get("answer") or "")
    evidence = case.get("evidence") if isinstance(case.get("evidence"), dict) else {}
    run_manifest = (
        evidence.get("run_manifest")
        if isinstance(evidence.get("run_manifest"), dict) else {}
    )
    if rating == "up" and len(answer) > 300:
        try:
            from hashmm.evolution.skill_manager import get_skill_manager
            mgr = get_skill_manager()
            if mgr.should_create_skill(query=query, answer=answer,
                                       feedback="up", strategy="grounded", owner_id=uid):
                mgr.create_from_conversation(query=query, answer=answer,
                                             llm_fn=app_state.llm_fn, owner_id=uid)
        except Exception as _e:
            log_suppressed(_obs_logger, _e)
    if rating in ("up", "down"):
        # Attribute the outcome to the exact skill prompt hash captured by the
        # server at answer time.  Client payloads cannot name a skill/version.
        try:
            from hashmm.evolution.skill_evolver import get_skill_evolver
            get_skill_evolver().record_run_feedback(
                run_manifest, rating, owner_id=uid,
            )
        except Exception as _e:
            log_suppressed(_obs_logger, _e)
        try:
            from hashmm.evolution.prompt_optimizer import get_prompt_optimizer
            get_prompt_optimizer().record_feedback(
                prompt_template="", task_type="", feedback=rating, query=query,
            )
        except Exception as _e:
            log_suppressed(_obs_logger, _e)
        try:
            from hashmm.evolution.user_model import get_user_model
            get_user_model().update_from_conversation(
                user_id=uid, query=query, answer=answer, feedback=rating,
            )
        except Exception as _e:
            log_suppressed(_obs_logger, _e)


@router.post("/conversations/{conv_id}/messages/{message_id}/review")
async def review_message_feedback(conv_id: str, message_id: str,
                                  req: FeedbackReviewRequest, request: Request):
    """V349: structured, owner-bound production feedback -> eval candidate.

    Client prose is never accepted as the query/answer. The backend captures
    both plus run/tool/retrieval evidence from the authoritative assistant row.
    Negative cases stay pending until an admin supplies a trusted reference.
    """
    require_conv_access(request, conv_id)
    uid = resolve_user_id(request)
    rating = str(req.rating or "").strip().lower()
    if rating not in ("up", "down", "clear"):
        raise HTTPException(400, "rating 仅支持 up、down、clear")
    try:
        case = db.record_message_feedback(
            user_id=uid, conv_id=conv_id, message_id=message_id,
            rating=rating, reason_code=req.reason_code, comment=req.comment,
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    if not case:
        # Same response for missing/non-owned message after conversation check.
        raise HTTPException(404, "消息不存在")
    _apply_feedback_learning(case, uid)
    db.audit(uid, (require_auth(request).get("sub") if uid != "anonymous" else "anonymous"),
             "message_feedback", f"{case.get('rating')}:{case.get('reason_code')}:{message_id}")
    return {"ok": True, "feedback_case": case,
            "reason_options": db.FEEDBACK_FAILURE_REASONS}


@router.post("/conversations/{conv_id}/messages/{msg_idx}/feedback")
async def message_feedback(conv_id: str, msg_idx: int, req: FeedbackRequest, request: Request):
    """Backward-compatible index endpoint using the same authoritative V349 path."""
    require_conv_access(request, conv_id)
    uid = resolve_user_id(request)
    feedback = str(req.feedback or "").lower()
    if feedback not in ("up", "down"):
        raise HTTPException(400, "feedback 仅支持 up 或 down")
    msgs = db.get_messages(conv_id, limit=10000)
    if msg_idx < 0 or msg_idx >= len(msgs) or msgs[msg_idx].get("role") != "assistant":
        raise HTTPException(404, "消息不存在")
    case = db.record_message_feedback(
        user_id=uid, conv_id=conv_id, message_id=str(msgs[msg_idx]["id"]),
        rating=feedback, reason_code="other" if feedback == "down" else "",
        comment="旧客户端未提供结构化失败原因" if feedback == "down" else "",
    )
    if not case:
        raise HTTPException(404, "消息不存在")
    _apply_feedback_learning(case, uid)
    return {"ok": True, "feedback_case": case}


@router.post("/{conv_id}/rewind", summary="回退到第 k 条 assistant 回复之前（含文件还原）")
async def rewind_conversation(conv_id: str, request: Request):
    """V273 CC 式回退：对话截断 + 会话工作区文件按快照还原（资料 13.3.5 两路线中的
    Claude Code 路线）。入参 assistant_index=从 0 计的 assistant 消息序号；服务端按
    created_at 排序取第 k 条 assistant 消息，删除它及之后的全部消息，并把工作区还原
    到它开始前的快照。快照缺失时只做对话截断并如实告知（files_restored=false）。"""
    require_auth(request)
    require_conv_access(request, conv_id)   # 归属校验：走既有统一入口
    body = await request.json()
    try:
        k = int(body.get("assistant_index"))
    except (TypeError, ValueError):
        raise HTTPException(400, "assistant_index 必填（从 0 计）")
    import sqlite3 as _sq
    from hashmm.api.database import DB_PATH
    conn = _sq.connect(str(DB_PATH))
    try:
        cur = conn.cursor()
        cols = [r[1] for r in cur.execute("PRAGMA table_info(messages)").fetchall()]
        conv_col = "conversation_id" if "conversation_id" in cols else ("conv_id" if "conv_id" in cols else None)
        ts_col = "created_at" if "created_at" in cols else ("ts" if "ts" in cols else None)
        if not conv_col or not ts_col:
            raise HTTPException(500, f"messages 表结构不识别：{cols[:6]}")
        rows = cur.execute(
            f"SELECT id, {ts_col} FROM messages WHERE {conv_col}=? AND role='assistant' ORDER BY {ts_col} ASC",
            (conv_id,)).fetchall()
        if k < 0 or k >= len(rows):
            raise HTTPException(404, f"assistant_index 越界（共 {len(rows)} 条）")
        target_id, target_ts = rows[k]
        cur.execute(f"DELETE FROM messages WHERE {conv_col}=? AND {ts_col}>=?", (conv_id, target_ts))
        removed = cur.rowcount
        cur.execute("DELETE FROM conversation_compactions WHERE conv_id=?", (conv_id,))
        conn.commit()
    finally:
        conn.close()
    from hashmm.api import conv_snapshots as snaps
    files_restored = snaps.restore(conv_id, str(target_id))
    try:
        db.audit(user["uid"], user.get("sub") or "", "rewind",
                 f"{conv_id} → assistant#{k}，删 {removed} 条，文件还原={files_restored}")
    except Exception:
        pass
    return {"ok": True, "removed": removed, "files_restored": files_restored,
            "note": "" if files_restored else "该节点无文件快照（多为老会话），仅回退了对话"}
