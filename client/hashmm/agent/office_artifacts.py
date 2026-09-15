"""Deterministic Office artifact inspection for Chat and Document Studio.

This is an in-process capability, not a wrapper around OfficeCLI.  It exposes a
small, stable contract over the Office libraries that HashMM already ships so
the Agent can inspect a deliverable before claiming that it is complete.
"""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any


SCHEMA = "hashmm.office-artifact.v1"
SUPPORTED = {".docx", ".pptx", ".xlsx"}
MAX_SCAN_CELLS = 50_000


def _base(path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    return {
        "schema": SCHEMA,
        "file": {
            "name": path.name,
            "format": path.suffix.lower().lstrip("."),
            "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
        },
        "verified": False,
        "structure": {},
        "warnings": [],
        "recommended_actions": [],
    }


def _inspect_docx(path: Path, report: dict[str, Any]) -> None:
    from docx import Document

    doc = Document(path)
    paragraphs = [p for p in doc.paragraphs if (p.text or "").strip()]
    headings = [p.text.strip() for p in paragraphs
                if str(getattr(p.style, "name", "")).lower().startswith("heading")]
    tables = []
    for table in doc.tables[:100]:
        tables.append({
            "rows": len(table.rows),
            "columns": max((len(r.cells) for r in table.rows), default=0),
        })
    image_relations = sum(
        1 for rel in doc.part.rels.values()
        if "image" in str(getattr(rel, "reltype", "")).lower()
    )
    report["structure"] = {
        "paragraphs": len(paragraphs),
        "headings": headings[:80],
        "heading_count": len(headings),
        "tables": tables,
        "table_count": len(doc.tables),
        "images": image_relations,
        "sections": len(doc.sections),
        "characters": sum(len(p.text or "") for p in paragraphs),
    }
    if not paragraphs and not doc.tables:
        report["warnings"].append("文档没有可读正文或表格")
    if len(paragraphs) >= 8 and not headings:
        report["warnings"].append("长文档没有标题层级，后续检索与阅读定位会较困难")
    if any(len(p.text or "") > 900 for p in paragraphs):
        report["warnings"].append("存在超过 900 字的长段落，建议拆分以改善阅读和 RAG 分块")
    report["recommended_actions"] = ["深度解读", "优化润色", "转换为结构化提要"]


def _inspect_pptx(path: Path, report: dict[str, Any]) -> None:
    from pptx import Presentation

    prs = Presentation(path)
    slides = []
    missing_titles = 0
    dense_slides = 0
    for number, slide in enumerate(prs.slides, 1):
        title = ""
        try:
            title = (slide.shapes.title.text or "").strip() if slide.shapes.title else ""
        except Exception:
            title = ""
        texts = []
        charts = tables = images = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = (getattr(shape, "text", "") or "").strip()
                if value:
                    texts.append(value)
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if int(getattr(shape, "shape_type", 0) or 0) == 13:  # MSO_SHAPE_TYPE.PICTURE
                images += 1
        chars = sum(len(value) for value in texts)
        if not title:
            missing_titles += 1
        if chars > 700 or len(texts) > 14:
            dense_slides += 1
        slides.append({"number": number, "title": title, "text_blocks": len(texts),
                       "characters": chars, "charts": charts, "tables": tables,
                       "images": images})
    report["structure"] = {
        "slides": slides,
        "slide_count": len(slides),
        "missing_title_slides": missing_titles,
        "dense_slides": dense_slides,
        "width": int(prs.slide_width),
        "height": int(prs.slide_height),
    }
    if not slides:
        report["warnings"].append("演示文稿没有幻灯片")
    if missing_titles:
        report["warnings"].append(f"{missing_titles} 页缺少可识别标题")
    if dense_slides:
        report["warnings"].append(f"{dense_slides} 页内容密度偏高，投屏阅读可能困难")
    report["recommended_actions"] = ["生成逐页摘要", "检查演示结构", "继续编辑幻灯片"]


def _inspect_xlsx(path: Path, report: dict[str, Any]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = []
        scanned = 0
        for ws in wb.worksheets:
            formulas = nonempty = 0
            for row in ws.iter_rows():
                for cell in row:
                    if scanned >= MAX_SCAN_CELLS:
                        break
                    scanned += 1
                    value = cell.value
                    if value is not None:
                        nonempty += 1
                    if isinstance(value, str) and value.startswith("="):
                        formulas += 1
                if scanned >= MAX_SCAN_CELLS:
                    break
            sheets.append({
                "name": ws.title,
                "rows": int(ws.max_row or 0),
                "columns": int(ws.max_column or 0),
                "nonempty_cells_scanned": nonempty,
                "formulas_scanned": formulas,
                "hidden": ws.sheet_state != "visible",
            })
        report["structure"] = {
            "sheets": sheets,
            "sheet_count": len(sheets),
            "cells_scanned": scanned,
            "scan_truncated": scanned >= MAX_SCAN_CELLS,
            "named_ranges": len(list(wb.defined_names.values())),
        }
    finally:
        wb.close()
    if not report["structure"]["sheets"]:
        report["warnings"].append("工作簿没有工作表")
    if report["structure"]["scan_truncated"]:
        report["warnings"].append("工作簿较大，本次结构检查只扫描前 5 万个单元格")
    if any(s["rows"] > 100_000 for s in report["structure"]["sheets"]):
        report["warnings"].append("存在超过 10 万行的工作表，建议使用分批分析")
    report["recommended_actions"] = ["数据问答", "生成统计摘要", "检查公式与数据口径"]


def inspect_office(path: str | Path) -> dict[str, Any]:
    """Parse and inspect one local Office artifact without invoking a model."""
    resolved = Path(path).resolve()
    if resolved.suffix.lower() not in SUPPORTED:
        raise ValueError("仅支持 .docx、.pptx、.xlsx")
    if not resolved.is_file():
        raise FileNotFoundError(resolved.name)
    report = _base(resolved)
    try:
        if resolved.suffix.lower() == ".docx":
            _inspect_docx(resolved, report)
        elif resolved.suffix.lower() == ".pptx":
            _inspect_pptx(resolved, report)
        else:
            _inspect_xlsx(resolved, report)
        report["verified"] = True
        report["integrity"] = "parsed"
    except Exception as exc:
        report["integrity"] = "invalid"
        report["warnings"].append(f"文件解析失败：{type(exc).__name__}")
    return report


def report_markdown(report: dict[str, Any]) -> str:
    """Render the bounded inspection contract as a user-facing audit report."""
    file = report.get("file") or {}
    structure = report.get("structure") or {}
    warnings = list(report.get("warnings") or [])
    lines = [
        f"# Office 结构检查：{file.get('name', '文件')}",
        "",
        f"- 格式：{str(file.get('format') or '').upper()}",
        f"- 大小：{int(file.get('size') or 0):,} 字节",
        f"- SHA-256：`{file.get('sha256', '')}`",
        f"- 可解析：{'是' if report.get('verified') else '否'}",
        "",
        "## 结构",
        "",
    ]
    fmt = file.get("format")
    if fmt == "docx":
        lines += [f"- {structure.get('paragraphs', 0)} 个正文段落，{structure.get('heading_count', 0)} 个标题",
                  f"- {structure.get('table_count', 0)} 个表格，{structure.get('images', 0)} 张图片，{structure.get('sections', 0)} 个分节"]
        for heading in list(structure.get("headings") or [])[:30]:
            lines.append(f"- 标题：{heading}")
    elif fmt == "pptx":
        lines += [f"- {structure.get('slide_count', 0)} 页幻灯片",
                  f"- {structure.get('missing_title_slides', 0)} 页缺少标题，{structure.get('dense_slides', 0)} 页密度偏高"]
        for slide in list(structure.get("slides") or [])[:20]:
            lines.append(f"- 第 {slide['number']} 页：{slide.get('title') or '未命名'}（{slide.get('text_blocks', 0)} 个文本块）")
    elif fmt == "xlsx":
        lines.append(f"- {structure.get('sheet_count', 0)} 个工作表，已扫描 {structure.get('cells_scanned', 0):,} 个单元格")
        for sheet in list(structure.get("sheets") or [])[:30]:
            lines.append(f"- {sheet['name']}：{sheet['rows']:,} 行 × {sheet['columns']:,} 列，扫描到 {sheet['formulas_scanned']:,} 个公式")
    lines += ["", "## 风险与建议", ""]
    lines += [f"- {warning}" for warning in warnings] or ["- 未发现结构性问题"]
    actions = list(report.get("recommended_actions") or [])
    if actions:
        lines += ["", "建议下一步：" + "、".join(actions) + "。"]
    lines += ["", "> 此报告只陈述程序实际解析到的结构；不把模型判断当作文件验证证据。"]
    return "\n".join(lines)


__all__ = ["SCHEMA", "SUPPORTED", "inspect_office", "report_markdown"]
